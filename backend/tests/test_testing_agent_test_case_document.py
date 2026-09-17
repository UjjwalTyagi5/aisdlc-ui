"""The Testing agent's test case document — the record, beside the spreadsheet.

The test plan was a bare pandas spreadsheet; the BRD and the design document it
sits beside on the approvals queue are designed documents. The test case document
is painted on the same shared canvas (`shared/docs/markdown_docx`): title band, facts
strip (what the cases were derived from), one numbered section per scenario type,
steps one per line, and — after a run — each executed test's outcome as a pill.
"""
from __future__ import annotations

from unittest.mock import AsyncMock, patch

import pytest
from docx import Document
from docx.oxml.ns import qn

from agents_orchestrator.testing_agent.config.session_state import TestCase, TestPlan
from agents_orchestrator.testing_agent.test_case_document import (
    TestDocMeta,
    _steps_cell,
    render_test_cases_docx,
)
from agents_orchestrator.testing_agent.test_case_document import test_cases_markdown as cases_markdown
from shared.docs import pwc_style as st

CASES = [
    TestCase(test_case_id="TC-REQ-01", feature_or_function_tested="Shorten a URL",
             test_summary="Valid intranet URL gets a slug", scenario_type="Happy Path",
             test_steps="1. Open /new 2. Paste https://intranet/policies 3. Submit",
             test_data="https://intranet/policies", expected_result="A slug is shown"),
    TestCase(test_case_id="TC-REQ-03", feature_or_function_tested="Shorten a URL",
             test_summary="External URL is refused", scenario_type="Error Case",
             test_steps="1. Paste https://example.com 2. Submit", test_data="https://example.com",
             expected_result="400 with 'intranet URLs only'"),
    TestCase(test_case_id="TC-REQ-04", feature_or_function_tested="Redirect",
             test_summary="Disabled slug | branded not-found", scenario_type="Edge Case",
             test_steps="Open /old-slug", test_data="slug=old-slug", expected_result="Branded 404"),
]


def _texts(doc) -> str:
    out = [p.text for p in doc.paragraphs]
    seen: list = []
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if any(cell._tc is tc for tc in seen):
                    continue
                seen.append(cell._tc)
                out.extend(p.text for p in cell.paragraphs)
    return "\n".join(out)


def test_the_markdown_groups_by_scenario_in_a_fixed_order_and_escapes_pipes():
    md = cases_markdown(CASES)
    assert md.index("## HAPPY PATH") < md.index("## ERROR CASES") < md.index("## EDGE CASES")
    assert "| TC-REQ-04 | Redirect | Disabled slug ∣ branded not-found |" in md, "a pipe must not split the row"
    assert "1. Open /new<br>2. Paste https://intranet/policies<br>3. Submit" in md


@pytest.mark.parametrize("raw, expected", [
    ("1. Open /new 2. Paste value 3. Submit", "1. Open /new<br>2. Paste value<br>3. Submit"),
    ("1) go 2) click", "1. go<br>2. click"),
    ("Open the page and submit", "Open the page and submit"),
    ("Precondition: logged in. 1. Open 2. Submit", "Precondition: logged in.<br>1. Open<br>2. Submit"),
    ("Enter 12 and 34", "Enter 12 and 34"),
])
def test_numbered_steps_go_one_per_line(raw, expected):
    assert _steps_cell(raw) == expected


def test_the_document_has_the_band_the_facts_and_the_sections(tmp_path):
    path = tmp_path / "test_cases.docx"
    render_test_cases_docx(CASES, str(path), meta=TestDocMeta(
        title="TEST Project — Test cases", project="TEST Project",
        source="the project's approved documents: QuickLink_BRD_new.docx, architecture.docx",
        generated_on="16 Sep 2026", test_types=["functional"],
    ))
    doc = Document(str(path))
    text = _texts(doc)
    assert "TEST CASE DOCUMENT" in text
    assert "3 test cases" in text
    assert "QuickLink_BRD_new.docx, architecture.docx" in text, "the fact names the documents"
    assert "the project's approved documents:" not in text, "…without the sentence's preamble"
    heads = [p.text for p in doc.paragraphs if p.text.startswith(("01", "02", "03"))]
    assert [h.split("  ", 1)[1] for h in heads] == ["HAPPY PATH", "ERROR CASES", "EDGE CASES"]
    assert doc.core_properties.title == "TEST Project — Test cases"


def test_results_paint_as_pills(tmp_path):
    path = tmp_path / "t.docx"
    render_test_cases_docx(CASES, str(path), meta=TestDocMeta(title="T"), results=[
        {"name": "TC-REQ-01", "status": "Pass", "detail": ""},
        {"name": "TC-REQ-03", "status": "Fail", "detail": "Got 500"},
        {"name": "TC-REQ-04", "status": "Blocked", "detail": "no fixture"},
    ])
    doc = Document(str(path))
    assert "1 of 3 passed" in _texts(doc)
    fills = {}
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        rpr = r._r.rPr
                        shd = rpr.find(qn("w:shd")) if rpr is not None else None
                        if shd is not None and r.text.strip() in ("Pass", "Fail", "Blocked"):
                            fills[r.text.strip()] = shd.get(qn("w:fill"))
    assert fills == {"Pass": st.GREEN_TINT, "Fail": st.RED_TINT, "Blocked": st.AMBER_TINT}


def test_an_empty_plan_still_renders_a_document(tmp_path):
    path = tmp_path / "empty.docx"
    render_test_cases_docx([], str(path), meta=TestDocMeta(title="Nothing"))
    assert "No test cases were generated" in _texts(Document(str(path)))


# ── the node writes it beside the spreadsheet ───────────────────────────────


async def test_package_final_reports_writes_the_document_and_announces_it(tmp_path, monkeypatch):
    from agents_orchestrator.testing_agent.Nodes import finalize as fz
    from agents_orchestrator.testing_agent.config import shared as ts

    announced = []

    async def _announce(session_id, filename, path):
        announced.append(filename)

    monkeypatch.setattr(fz, "BROADCASTING_AVAILABLE", True)
    monkeypatch.setattr(fz, "get_session_id", lambda: "sess-1")
    monkeypatch.setattr(fz, "get_user_id", lambda: "u1")
    monkeypatch.setattr(fz, "broadcast_file_generated", _announce)
    monkeypatch.setattr(fz.esett, "FILES", str(tmp_path))

    state = {
        "test_plan": TestPlan(test_cases=CASES),
        "project_display_name": "TEST Project",
        "plan_source": "the project's approved documents: QuickLink_BRD_new.docx",
        "selected_test_types": ["functional"],
        "ui_test_results": [{"id": "TC-REQ-01", "description": "x", "status": "Pass", "detail": ""}],
    }
    out = await fz.package_final_reports(state)

    out_dir = tmp_path / "u1" / "orchestrator" / "sess-1" / "output"
    assert (out_dir / "test_plan.xlsx").exists() and (out_dir / "test_cases.docx").exists()
    assert announced[:2] == ["test_plan.xlsx", "test_cases.docx"]
    assert "excel_plan_b64" in out["final_outputs"]
    text = _texts(Document(str(out_dir / "test_cases.docx")))
    assert "TEST Project — Test cases" in text and "1 of 1 passed" in text
