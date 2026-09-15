"""ADRs paint as cards and judgement words as pills — the palette's rule made visible.

The brief's rule: red = risk, amber = needs confirming, green only "fine as is", grey =
context. In a design document the words that carry those judgements sit in tables —
a security checklist's "Risk" column, an ADR's "Status" — and painted as body text
they read no differently from the words around them. So a High risk is a red pill,
an Accepted decision a green one, and each `#### ADR-nnn` opens a card with an orange
edge that holds the decision's fields as label-over-value, so one decision reads as
one unit rather than a heading floating over a two-column table.

The column matters: "High" under "Talks to" is a name, not a verdict.
"""
from __future__ import annotations

import pytest
from docx import Document
from docx.oxml.ns import qn

from agents_orchestrator.design_architecture_agent.design_document import (
    DesignMeta,
    pill_for,
    render_design_docx,
)
from shared.docs import pwc_style as st

SAMPLE = """\
# Payments Gateway

## ARCHITECTURE DECISION RECORDS

#### ADR-001: Use PostgreSQL for the ledger

| Field | Details |
|-------|---------|
| Status | Accepted |
| Context | The ledger needs strict consistency |
| Decision | PostgreSQL 16 with `SERIALIZABLE` isolation |
| Alternatives | DynamoDB — rejected for lack of multi-row transactions |

#### ADR-002: Event log between modules

| Field | Details |
|-------|---------|
| Status | Proposed |
| Context | Fulfilment will be extracted later |

## SECURITY DESIGN CHECKLIST

| OWASP Category | Applicable? | Risk in this design | Mitigation control |
|----------------|-------------|---------------------|--------------------|
| Injection | Yes | High | Parameterised queries only |
| Logging failures | Yes | Low | Structured audit log |

| Service | Talks to |
|---------|----------|
| Ledger | High |
"""


@pytest.fixture
def doc(tmp_path):
    path = tmp_path / "adr.docx"
    render_design_docx(SAMPLE, str(path), meta=DesignMeta(title="Payments Gateway", components=["adr", "security"]))
    return Document(str(path))


def _cells(doc):
    seen: list = []
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if any(cell._tc is tc for tc in seen):
                    continue
                seen.append(cell._tc)
                yield cell


def _runs(doc):
    for p in doc.paragraphs:
        yield from p.runs
    for cell in _cells(doc):
        for p in cell.paragraphs:
            yield from p.runs


def _shading(run):
    rpr = run._r.rPr
    shd = rpr.find(qn("w:shd")) if rpr is not None else None
    return shd.get(qn("w:fill")) if shd is not None else None


def test_each_adr_is_a_card_with_an_orange_edge_holding_its_fields(doc):
    cards = [c for c in _cells(doc) if "ADR-001" in c.text]
    assert len(cards) == 1, "one card per decision"
    card = cards[0]
    borders = card._tc.tcPr.find(qn("w:tcBorders"))
    assert borders is not None and borders.find(qn("w:left")).get(qn("w:color")) == st.ACCENT
    text = card.text
    for field in ("Status", "Context", "Decision", "Alternatives"):
        assert field in text, field  # painted in caps via the run property; text unchanged
    assert "The ledger needs strict consistency" in text
    assert "Use PostgreSQL for the ledger" in text
    # the field table was NOT painted as a second, free-standing table
    assert not any(c.text.strip() == "Field" for c in _cells(doc))


def test_the_second_adr_gets_its_own_card(doc):
    assert any("ADR-002" in c.text and "Fulfilment" in c.text for c in _cells(doc))
    assert not any("ADR-001" in c.text and "ADR-002" in c.text for c in _cells(doc))


def test_an_accepted_status_is_a_green_pill_and_proposed_is_amber(doc):
    fills = {r.text.strip(): _shading(r) for r in _runs(doc) if r.text.strip() in ("Accepted", "Proposed")}
    assert fills["Accepted"] == st.GREEN_TINT
    assert fills["Proposed"] == st.AMBER_TINT


def test_a_high_risk_is_a_red_pill_and_a_low_one_is_grey(doc):
    highs = [r for r in _runs(doc) if r.text.strip() == "High"]
    lows = [r for r in _runs(doc) if r.text.strip() == "Low"]
    assert lows and _shading(lows[0]) == st.GREY_TINT
    # two "High" cells: the risk one is a pill, the "Talks to" one is a plain name
    assert {str(_shading(r)) for r in highs} == {"None", st.RED_TINT}


def test_the_section_after_an_adr_is_not_swallowed_by_its_card(doc):
    heads = [p.text for p in doc.paragraphs if p.text.startswith(("01", "02"))]
    assert len(heads) == 2 and "Security" in heads[1]
    assert not any("OWASP" in c.text and "ADR-002" in c.text for c in _cells(doc))


@pytest.mark.parametrize("value, column, expected", [
    ("High", "Risk in this design", (st.RED_TINT, st.RED)),
    ("**Critical**", "Severity", (st.RED_TINT, st.RED)),
    ("medium", "Priority", (st.AMBER_TINT, st.ACCENT_DEEP)),
    ("Low", "Impact", (st.GREY_TINT, st.GREY_INK)),
    ("Accepted", "Status", (st.GREEN_TINT, st.GREEN)),
    ("Superseded", "Status", (st.GREY_TINT, st.GREY_INK)),
    ("High", "Talks to", None),
    ("High availability", "Risk", None),
    ("Yes", "Applicable?", None),
])
def test_pill_for(value, column, expected):
    assert pill_for(value, column=column) == expected
