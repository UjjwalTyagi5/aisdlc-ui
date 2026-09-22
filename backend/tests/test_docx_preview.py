"""Reading a stored Word document back as markdown, for documents with no page copy.

LIVE (2026-09-19): clicking a generated design showed a download card and the file name —
`architecture.docx` had no page copy, because the Design agent writes no sibling markdown.
The two things the real file broke are pinned here: it names no heading styles (every
paragraph is "Normal", the design is direct formatting) and its title band is one merged
cell that python-docx repeats once per column it spans — 24 times.
"""
from __future__ import annotations

import base64
import io
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from shared.services.docx_preview import can_preview, docx_markdown  # noqa: E402


def _bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _designed_document():
    """A document shaped like the agents': no heading styles, a merged band, tables."""
    d = Document()

    band = d.add_table(rows=1, cols=4)
    band.rows[0].cells[0].merge(band.rows[0].cells[3]).text = "DESIGN DOCUMENT · Track 1"

    for text, size in (("01  Overview", 12.5), ("Executive Summary", 10.5)):
        p = d.add_paragraph()
        run = p.add_run(text)
        run.bold, run.font.size = True, Pt(size)

    body = d.add_paragraph()
    body.add_run("QuickLink replaces a withdrawn public shortener.").font.size = Pt(9)

    table = d.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Field", "Details"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Objective", "Track clicks | safely"
    return d


def test_the_designed_document_reads_as_headings_body_and_tables():
    markdown = docx_markdown(_bytes(_designed_document()))
    lines = markdown.splitlines()

    # The band is one line, not the same cell 4 times.
    assert lines[0] == "DESIGN DOCUMENT · Track 1"
    assert markdown.count("DESIGN DOCUMENT") == 1

    # Headings come from size and weight, ranked — there are no heading STYLES here.
    assert "## 01  Overview" in markdown
    assert "### Executive Summary" in markdown
    assert "QuickLink replaces a withdrawn public shortener." in markdown

    assert "| Field | Details |" in markdown
    assert "| --- | --- |" in markdown
    # A pipe inside a cell would otherwise split the column.
    assert "| Objective | Track clicks \\| safely |" in markdown


def test_named_heading_styles_are_used_when_the_document_has_them():
    d = Document()
    d.add_heading("Architecture", level=1)
    d.add_heading("Data flow", level=2)
    d.add_paragraph("Plain body text.")
    d.add_paragraph("First point", style="List Bullet")
    markdown = docx_markdown(_bytes(d))
    assert "# Architecture" in markdown
    assert "## Data flow" in markdown
    assert "- First point" in markdown


def test_an_empty_or_unreadable_file_asks_for_no_preview():
    assert docx_markdown(b"not a word file at all") is None
    assert docx_markdown(_bytes(Document())) is None


def test_only_word_documents_are_previewed():
    assert can_preview("architecture.docx")
    assert not can_preview("QuickLink_Project_Plan.xlsx")
    assert not can_preview("QuickLink_Architecture_Presentation.pptx")


def test_a_very_long_document_is_cut_and_says_so(monkeypatch):
    monkeypatch.setattr("shared.services.docx_preview.MAX_CHARS", 400)
    d = Document()
    for i in range(40):
        d.add_paragraph(f"Paragraph {i} with enough words to push past the limit quickly.")
    markdown = docx_markdown(_bytes(d))
    assert len(markdown) < 700
    assert markdown.endswith("_The rest of this document is in the Word file._")


@pytest.mark.parametrize("name", ["architecture.docx", "ARCHITECTURE.DOCX"])
def test_the_extension_check_ignores_case(name):
    assert can_preview(name)


# ── figures ──────────────────────────────────────────────────────────────────
# LIVE (2026-09-21): the Design agent's C4, sequence and ER diagrams are pictures in the Word
# file (`WordCanvas.figure`), and a design read back without a page copy showed only their
# captions — "Figure 3 · …" under nothing.

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)


def _with_figure():
    d = Document()
    d.add_paragraph("Context")
    d.add_picture(io.BytesIO(_PNG))
    d.add_paragraph("Figure 1  ·  System context")
    return d


def test_a_diagram_comes_through_in_place_above_its_caption():
    markdown = docx_markdown(_bytes(_with_figure()))
    figure = "![Figure](data:image/png;base64," + base64.b64encode(_PNG).decode("ascii") + ")"
    assert figure in markdown
    assert markdown.index("Context") < markdown.index(figure) < markdown.index("Figure 1  ·  System context")


def test_a_figure_past_the_page_budget_is_named_not_dropped(monkeypatch):
    monkeypatch.setattr("shared.services.docx_preview.MAX_FIGURES_BYTES", 10)
    markdown = docx_markdown(_bytes(_with_figure()))
    assert "data:image" not in markdown
    assert "_A figure here is only in the Word file._" in markdown


def test_figures_do_not_count_against_the_text_limit(monkeypatch):
    monkeypatch.setattr("shared.services.docx_preview.MAX_CHARS", 60)
    markdown = docx_markdown(_bytes(_with_figure()))
    assert "data:image/png" in markdown and "System context" in markdown
