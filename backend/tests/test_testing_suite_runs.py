"""Running suites: every case accounted for, and each status meaning what it says."""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from agents_orchestrator.testing_agent.suites import api_run, functional_run, reports, unit_run  # noqa: E402
from agents_orchestrator.testing_agent.suites.models import validate_cases  # noqa: E402


async def _quiet(_msg: str) -> None:
    return None


# ── API ───────────────────────────────────────────────────────────────────────

API_CASES = [
    {"id": "AT-001", "title": "Create", "method": "POST", "path": "/api/shorten", "body": {"longUrl": "https://x"},
     "expected_status": 201, "expected_body_contains": ["shortCode"], "capture": {"code": "shortCode"}},
    {"id": "AT-002", "title": "Stats", "method": "GET", "path": "/api/links/{{code}}/stats", "expected_status": 200,
     "expected_body_contains": ["totalClicks"]},
    {"id": "AT-003", "title": "Bad URL", "method": "POST", "path": "/api/shorten", "body": {"longUrl": "ftp://x"},
     "expected_status": 400, "expected_body_contains": ["Invalid URL"]},
]


def _client_with(handler, monkeypatch):
    real = httpx.AsyncClient

    def factory(*args, **kwargs):
        kwargs["transport"] = httpx.MockTransport(handler)
        return real(*args, **kwargs)

    monkeypatch.setattr(api_run.httpx, "AsyncClient", factory)


async def test_api_cases_run_in_order_with_captured_values(monkeypatch):
    seen = []

    def handler(request: httpx.Request):
        seen.append(f"{request.method} {request.url.path}")
        if request.url.path == "/api/shorten":
            body = json.loads(request.content)
            if body["longUrl"].startswith("ftp"):
                return httpx.Response(400, json={"error": "Invalid URL"})
            return httpx.Response(201, json={"shortCode": "abc1234"})
        if request.url.path == "/api/links/abc1234/stats":
            return httpx.Response(500, text="boom")
        return httpx.Response(200, text="ok")

    _client_with(handler, monkeypatch)
    cases, _ = validate_cases("api", API_CASES)
    rows, responses = await api_run.execute(cases, "http://app:8080", _quiet)
    assert "GET /api/links/abc1234/stats" in seen
    assert [r.status for r in rows] == ["Passed", "Failed", "Passed"]
    assert "Expected status 200, got 500" in rows[1].message and '"totalClicks"' in rows[1].message
    assert responses[1][2] == 500


async def test_a_case_needing_a_value_that_was_never_captured_is_not_run(monkeypatch):
    def handler(request: httpx.Request):
        if request.url.path == "/api/shorten":
            return httpx.Response(500, json={"error": "db locked"})
        return httpx.Response(200, text="ok")

    _client_with(handler, monkeypatch)
    cases, _ = validate_cases("api", API_CASES[:2])
    rows, _ = await api_run.execute(cases, "http://app:8080", _quiet)
    assert rows[0].status == "Failed"
    assert rows[1].status == "Not run" and "{{code}}" in rows[1].message and "AT-001" in rows[1].message


async def test_an_app_that_goes_down_mid_run_is_said_once_and_the_rest_are_not_run(monkeypatch):
    """LIVE: the app stopped during the first request and seven cases each said 'ConnectError'."""
    state = {"up": True}

    def handler(request: httpx.Request):
        if not state["up"]:
            raise httpx.ConnectError("refused", request=request)
        if request.url.path == "/api/shorten":
            state["up"] = False
            raise httpx.ReadError("connection closed", request=request)
        return httpx.Response(200, text="ok")

    _client_with(handler, monkeypatch)
    cases, _ = validate_cases("api", API_CASES)
    rows, _ = await api_run.execute(cases, "http://app:8080", _quiet)
    assert rows[0].status == "Error" and "stopped responding" in rows[0].message
    assert [r.status for r in rows[1:]] == ["Not run", "Not run"]
    assert all("during AT-001" in r.message for r in rows[1:])


async def test_an_app_that_is_not_running_is_refused_before_any_case(monkeypatch):
    def handler(request: httpx.Request):
        raise httpx.ConnectError("refused", request=request)

    _client_with(handler, monkeypatch)
    cases, _ = validate_cases("api", API_CASES)
    with pytest.raises(RuntimeError, match="did not respond"):
        await api_run.execute(cases, "http://app:8080", _quiet)


# ── unit ──────────────────────────────────────────────────────────────────────


def test_case_ids_are_read_from_jest_and_pytest_names():
    assert unit_run.canonical_id("UT-012: rejects an empty URL") == "UT-012"
    assert unit_run.canonical_id("urlService › UT-003: x") == "UT-003"
    assert unit_run.canonical_id("test_UT_007_disables") == "UT-007"
    assert unit_run.canonical_id("UT-0071: x") == "UT-0071"
    assert unit_run.canonical_id("no id here") is None


async def test_jest_results_map_to_cases_and_a_file_that_did_not_load_is_an_error(tmp_path, monkeypatch):
    gen = tmp_path / unit_run.GEN_DIR
    gen.mkdir()
    ok_file, broken_file = gen / "src_a.test.js", gen / "src_b.test.js"
    ok_file.write_text("test('UT-001: a', () => {})")
    broken_file.write_text("require('../nope')")

    async def fake_run(cmd, cwd, timeout, env=None):
        out = next(a.split("=", 1)[1] for a in cmd if a.startswith("--outputFile="))
        Path(out).write_text(json.dumps({"testResults": [
            {"name": str(ok_file), "status": "failed", "message": "", "assertionResults": [
                {"title": "UT-001: a", "status": "passed", "duration": 4, "failureMessages": []},
                {"title": "UT-002: b", "status": "failed", "duration": 2, "failureMessages": ["\x1b[31mExpected 7, got 6\x1b[0m"]},
            ]},
            {"name": str(broken_file), "status": "failed", "message": "Cannot find module '../nope'", "assertionResults": []},
        ]}))
        return type("R", (), {"timed_out": False, "stdout": "", "stderr": "", "ok": True})()

    monkeypatch.setattr(unit_run, "_run", fake_run)
    files, results = await unit_run.run_node(str(tmp_path), ["node", "jest.js"], [f"{unit_run.GEN_DIR}/src_a.test.js", f"{unit_run.GEN_DIR}/src_b.test.js"])
    assert results["UT-001"]["status"] == "Passed" and results["UT-002"]["status"] == "Failed"
    assert results["UT-002"]["message"] == "Expected 7, got 6"
    assert files[f"{unit_run.GEN_DIR}/src_b.test.js"].loaded is False
    assert "Cannot find module" in files[f"{unit_run.GEN_DIR}/src_b.test.js"].message


def test_a_repository_that_is_neither_node_nor_python_is_refused(tmp_path):
    with pytest.raises(unit_run.UnitRunError, match="Node.js \\(Jest\\) and Python \\(pytest\\)"):
        unit_run.detect_stack(str(tmp_path))


# ── functional ────────────────────────────────────────────────────────────────


class _FakeBrowser:
    def __init__(self, fail_at=None, error=None):
        self.fail_at, self.error, self.ran = fail_at, error, []
        self.d = type("D", (), {"delete_all_cookies": lambda self: None, "save_screenshot": lambda self, p: True})()

    def run_step(self, step):
        self.ran.append(step.action)
        if len(self.ran) == self.fail_at:
            raise self.error
        return "ok"


def _func_case():
    cases, _ = validate_cases("functional", [{"id": "FT-001", "title": "Create", "expected": "Created", "steps": [
        {"action": "open", "target": "/"}, {"action": "type", "target": "Long URL", "value": "https://x"},
        {"action": "click", "target": "Create"}, {"action": "assert_text", "value": "Created"}]}])
    return cases[0]


def test_a_journey_that_passes_every_step_passes(tmp_path, monkeypatch):
    monkeypatch.setattr(functional_run.time, "sleep", lambda _s: None)
    row, steps = functional_run.run_case(_FakeBrowser(), _func_case(), str(tmp_path))
    assert row.status == "Passed" and [s[3] for s in steps] == ["Passed"] * 4


def test_a_check_that_does_not_hold_fails_the_case_and_stops_it(tmp_path, monkeypatch):
    monkeypatch.setattr(functional_run.time, "sleep", lambda _s: None)
    browser = _FakeBrowser(fail_at=4, error=functional_run.StepFailure('the page does not show "Created"'))
    row, steps = functional_run.run_case(browser, _func_case(), str(tmp_path))
    assert row.status == "Failed" and row.message.startswith("Step 4 (Check the page shows")
    assert row.evidence == "FT-001_step4.png"


def test_an_unreachable_page_is_an_error_and_later_steps_are_not_run(tmp_path, monkeypatch):
    monkeypatch.setattr(functional_run.time, "sleep", lambda _s: None)
    browser = _FakeBrowser(fail_at=1, error=functional_run.PageError("http://app/ could not be reached"))
    row, steps = functional_run.run_case(browser, _func_case(), str(tmp_path))
    assert row.status == "Error" and browser.ran == ["open"]
    assert [s[3] for s in steps] == ["Failed", "Not run", "Not run", "Not run"]


# ── the report ────────────────────────────────────────────────────────────────


def test_the_report_accounts_for_every_case_with_a_verdict():
    rows = [reports.ResultRow("UT-001", "a", "m", "Passed", 3), reports.ResultRow("UT-002", "b", "m", "Not run", None, "why")]
    meta = reports.ReportMeta(kind="unit", repository="QuickLink", branch="main", commit="082f91e49bf0",
                              suite_document="QuickLink_Unit_Test_Cases.xlsx", suite_status="approved", runner="Jest (Node.js)")
    assert reports.verdict(rows) == "Incomplete"
    wb = load_workbook(io.BytesIO(reports.write_report(meta, rows, {"Generated tests": ([("File", 10), ("Code", 10)], [["f", "x" * 40_000]])})))
    assert wb.sheetnames == ["Summary", "Results", "Generated tests"]
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=3, values_only=True) if r[0]}
    assert summary["Result"] == "Incomplete" and summary["Pass rate"] == "50%"
    assert summary["Test cases"] == "QuickLink_Unit_Test_Cases.xlsx (approved)"
    assert [r[3] for r in wb["Results"].iter_rows(min_row=2, values_only=True)] == ["Passed", "Not run"]
    assert len(wb["Generated tests"]["B2"].value) < 32_767
    md = reports.report_markdown(meta, rows)
    assert "**Incomplete** — 1 passed, 0 failed, 0 could not run, 1 not run, of 2 cases." in md
