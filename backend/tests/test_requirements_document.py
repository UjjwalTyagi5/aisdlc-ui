"""The Requirements agent's BRD/PDD/risk register on the platform's designed canvas.

The BRD was written by a generic markdown→Word walker: Word's own "Heading 1",
"Table Grid", no title band, no facts. Now it paints on `shared/docs/markdown_docx`
like the brief, the design document and the test case document, and leaves its
markdown beside the .docx so the Requirements page can render it as a report.
"""
from __future__ import annotations

import os

import pytest
from docx import Document

from agents_orchestrator.requirements_agent import requirements_document as rd

BRD = """# QuickLink — Business Requirements

## Executive Summary
QuickLink shortens URLs for the marketing team.

## Project Objectives
- Cut link length by 80%
- Track clicks per campaign

## Project Scope
| In scope | Out of scope |
|---|---|
| Shortening | Billing |

## Requirements
1. Create a short link
2. View click counts

## Key Stakeholders
Marketing lead, platform team.
"""

PDD = """## Process Overview
Claims come in by email.

## As-Is Process
1. Triage
2. Assign

## To-Be Process
1. Auto-triage
"""

RISK = """## Risk Register
| Risk ID | Risk Description | Likelihood | Mitigation |
|---|---|---|---|
| R1 | Vendor delay | Medium | Second vendor |
"""


@pytest.mark.parametrize("markdown, filename, expected", [
    (BRD, "whatever.docx", "brd"),
    (PDD, "whatever.docx", "pdd"),
    (RISK, "whatever.docx", "risk"),
    ("## Notes\nhello", "QuickLink_BRD.docx", "brd"),
    ("## Notes\nhello", "process_definition.docx", "pdd"),
    ("## Notes\nhello", "risk_register_v2.docx", "risk"),
    ("## Notes\nhello", "output.docx", "document"),
])
def test_the_kind_is_read_off_the_headers_then_the_file_name(markdown, filename, expected):
    assert rd.document_kind(markdown, filename).id == expected


def test_one_header_hit_is_not_enough_to_call_it_a_brd():
    assert rd.document_kind("## Requirements\n- one", "notes.docx").id == "document"


def test_the_title_is_the_h1_else_project_and_kind_else_the_file_name():
    brd = rd.document_kind(BRD)
    assert rd.title_for(BRD, "x.docx", brd, "QuickLink") == "QuickLink — Business Requirements"
    assert rd.title_for(PDD, "x.docx", rd.document_kind(PDD), "ClaimTrack") == "ClaimTrack — Process definition"
    assert rd.title_for(PDD, "claims_process_v2.docx", rd.document_kind(PDD), "") == "claims process v2"


def test_the_word_file_carries_the_title_band_facts_and_sections(tmp_path):
    path = str(tmp_path / "QuickLink_BRD.docx")
    meta = rd.RequirementsDocMeta(
        title="QuickLink — Business Requirements", kind=rd.document_kind(BRD),
        project="QuickLink", sources=["brief.pdf", "notes.txt"], generated_on="16 Sep 2026",
    )
    assert rd.render_requirements_docx(BRD, path, meta=meta) == path
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    band = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "BUSINESS REQUIREMENTS DOCUMENT" in band and "QuickLink — Business Requirements" in band, "the title band"
    assert "Executive Summary" in text and "Key Stakeholders" in text
    assert doc.core_properties.title == "QuickLink — Business Requirements"
    assert "brief.pdf, notes.txt" in band and "16 Sep 2026" in band and "5" in band, "Source, Generated, Sections"
    assert "Billing" in band, "the scope table is painted"
    assert not any(p.style.name == "Heading 1" and p.text == "Executive Summary" for p in doc.paragraphs), \
        "sections are numbered on the canvas, not Word's default Heading 1"


def test_the_markdown_is_left_beside_the_docx_for_the_page(tmp_path):
    docx_path = str(tmp_path / "QuickLink_BRD.docx")
    md_path = rd.write_markdown_sibling("\r\n" + BRD + "\r\n", docx_path)
    assert md_path == str(tmp_path / "QuickLink_BRD.md")
    with open(md_path, encoding="utf-8") as fh:
        body = fh.read()
    assert body.startswith("# QuickLink") and body.endswith("platform team.\n") and "\r" not in body


def test_prompt_numbering_is_dropped_and_dotted_headers_become_subheads():
    src = "##1. Executive Summary\n## 1.1 Introduction\n## 3.6 Risks (Table)\n### 2 Data Flow\n## Glossary"
    assert rd.normalise_headers(src).splitlines() == [
        "## Executive Summary", "### Introduction", "### Risks (Table)", "### Data Flow", "## Glossary",
    ]
    pdd = "##1 Executive Summary\n##2 Solution Design Details\n##3 Scope of Requirements\n## 3.5 Assumptions"
    assert rd.document_kind(pdd).id == "pdd"


def test_source_names_are_base_names_and_skip_blanks():
    assert rd.source_names(["C:\\up\\brief.pdf", "", "/tmp/notes.txt"]) == ["brief.pdf", "notes.txt"]
