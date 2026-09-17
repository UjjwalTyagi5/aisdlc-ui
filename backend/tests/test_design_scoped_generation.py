"""The Design agent generates the components it was asked for, and adds to a document.

`test_design_components.py` proves the catalogue and the prompt builder. This file
proves the TOOLS use them: what reaches the model, what the session's document becomes
after a second request, and what the agent is told in its system prompt.

Stubbed: the model call (the prompt it receives is what is pinned), the file write and
the artifact registration — as the rest of this directory stubs them.
"""
from __future__ import annotations

from unittest.mock import AsyncMock, MagicMock, patch

import pytest

RUN_ID = "11111111-2222-3333-4444-555555555555"
USER_ID = "af57932d-f4f2-4673-aef7-d33b75c022f4"

HLD_DOC = "# QuickLink\n\n## HIGH-LEVEL DESIGN\n\nThe HLD body.\n"
DB_DOC = "## DATABASE SCHEMA\n\n```sql\nCREATE TABLE link (id INT);\n```\n"
DB_DOC_V2 = "## DATABASE SCHEMA\n\n```sql\nCREATE TABLE link (id BIGINT);\n```\n"


@pytest.fixture
def design(tmp_path):
    from agents_orchestrator.design_architecture_agent.agents import architecture as a
    from agents_orchestrator.design_architecture_agent.config import shared
    from shared.services import chat_artifacts

    original_dir = a._FILES_DIR
    a._FILES_DIR = str(tmp_path)
    shared.last_architecture = ""
    shared.output_file = ""
    shared.output_file_url = ""
    fake_manager = MagicMock()
    fake_manager.broadcast = AsyncMock()
    llm = AsyncMock(return_value=HLD_DOC)
    try:
        with patch.object(a, "get_user_id", lambda: USER_ID), \
                patch.object(a, "get_session_id", lambda: RUN_ID), \
                patch.object(a, "manager", fake_manager), \
                patch.object(a, "_llm_generate_async", llm), \
                patch.object(a, "_autosave_architecture", AsyncMock(return_value="SAVED: x.docx")), \
                patch.object(chat_artifacts, "register_generated_file", AsyncMock()):
            yield a, shared, llm
    finally:
        a._FILES_DIR = original_dir
        shared.last_architecture = ""
        shared.output_file = ""
        shared.output_file_url = ""


def _prompt_sent(llm) -> str:
    args, kwargs = llm.call_args
    return args[0] if args else kwargs.get("prompt", "")


# ── what reaches the model ───────────────────────────────────────────────────


async def test_the_model_is_asked_for_only_the_requested_components(design):
    a, shared, llm = design

    await a.generate_architecture_from_context.ainvoke(
        {"context": "QuickLink BRD text", "components": ["hld"]},
    )

    prompt = _prompt_sent(llm)
    assert "## HIGH-LEVEL DESIGN" in prompt
    assert "## DATABASE SCHEMA" not in prompt
    assert "## SECURITY DESIGN CHECKLIST" not in prompt
    assert "QuickLink BRD text" in prompt


async def test_the_upload_entry_point_scopes_the_same_way(design):
    a, shared, llm = design

    await a.generate_architecture.ainvoke(
        {"document_text": "uploaded requirements", "components": ["db", "api"]},
    )

    prompt = _prompt_sent(llm)
    assert "## DATABASE SCHEMA" in prompt and "## API CONTRACT" in prompt
    assert "## HIGH-LEVEL DESIGN" not in prompt


async def test_an_unknown_component_is_refused_without_calling_the_model(design):
    a, shared, llm = design

    out = await a.generate_architecture_from_context.ainvoke(
        {"context": "x", "components": ["deployment runbook"]},
    )

    assert "deployment runbook" in out
    assert "High-level design" in out, "the refusal lists what can be produced"
    llm.assert_not_called()
    assert shared.last_architecture == ""


# ── the session's document grows ─────────────────────────────────────────────


async def test_a_second_request_adds_to_the_document_in_catalogue_order(design):
    a, shared, llm = design
    await a.generate_architecture_from_context.ainvoke({"context": "ctx", "components": ["hld"]})

    llm.return_value = DB_DOC
    out = await a.generate_architecture_from_context.ainvoke({"context": "ctx", "components": ["db"]})

    doc = shared.last_architecture
    assert "## HIGH-LEVEL DESIGN" in doc and "## DATABASE SCHEMA" in doc
    assert doc.index("## HIGH-LEVEL DESIGN") < doc.index("## DATABASE SCHEMA")
    assert "The HLD body." in doc, "the earlier section must survive untouched"
    assert "## DATABASE SCHEMA" in out and "## HIGH-LEVEL DESIGN" in out


async def test_regenerating_a_component_replaces_only_that_section(design):
    a, shared, llm = design
    await a.generate_architecture_from_context.ainvoke({"context": "ctx", "components": ["hld"]})
    llm.return_value = DB_DOC
    await a.generate_architecture_from_context.ainvoke({"context": "ctx", "components": ["db"]})

    llm.return_value = DB_DOC_V2
    await a.generate_architecture_from_context.ainvoke({"context": "ctx", "components": ["db"]})

    doc = shared.last_architecture
    assert doc.count("## DATABASE SCHEMA") == 1
    assert "BIGINT" in doc and "id INT)" not in doc
    assert "The HLD body." in doc


async def test_the_model_is_told_what_the_document_already_holds(design):
    """Consistency across components: a DB schema written blind to the HLD names
    different entities. The existing sections go to the model as context."""
    a, shared, llm = design
    await a.generate_architecture_from_context.ainvoke({"context": "ctx", "components": ["hld"]})

    llm.return_value = DB_DOC
    await a.generate_architecture_from_context.ainvoke({"context": "ctx", "components": ["db"]})

    prompt = _prompt_sent(llm)
    assert "The HLD body." in prompt


def test_the_document_title_is_kept_across_additions():
    from agents_orchestrator.design_architecture_agent import components as dc

    merged = dc.merge_sections(HLD_DOC, DB_DOC)
    assert merged.startswith("# QuickLink")
    assert merged.index("## HIGH-LEVEL DESIGN") < merged.index("## DATABASE SCHEMA")


def test_merge_orders_by_catalogue_not_arrival():
    from agents_orchestrator.design_architecture_agent import components as dc

    merged = dc.merge_sections(DB_DOC, HLD_DOC)
    assert merged.index("## HIGH-LEVEL DESIGN") < merged.index("## DATABASE SCHEMA")


def test_sections_present_lists_what_the_document_holds():
    from agents_orchestrator.design_architecture_agent import components as dc

    assert dc.sections_present(HLD_DOC + DB_DOC) == ["hld", "db"]
    assert dc.sections_present("") == []


# ── the file the agent writes ────────────────────────────────────────────────


@pytest.fixture
def design_writing(tmp_path):
    """Like `design`, but the save path is REAL: the .docx is written under tmp_path."""
    from agents_orchestrator.design_architecture_agent.agents import architecture as a
    from agents_orchestrator.design_architecture_agent.config import shared
    from shared.services import chat_artifacts

    original_dir = a._FILES_DIR
    a._FILES_DIR = str(tmp_path)
    shared.last_architecture = ""
    shared.output_file = ""
    shared.output_file_url = ""
    fake_manager = MagicMock()
    fake_manager.broadcast = AsyncMock()
    llm = AsyncMock(return_value=HLD_DOC)
    try:
        with patch.object(a, "get_user_id", lambda: USER_ID), \
                patch.object(a, "get_session_id", lambda: RUN_ID), \
                patch.object(a, "manager", fake_manager), \
                patch.object(a, "_llm_generate_async", llm), \
                patch.object(a, "_render_mermaid_to_png", lambda code: None), \
                patch.object(chat_artifacts, "register_generated_file", AsyncMock()):
            yield a, shared, llm, tmp_path / USER_ID / "orchestrator" / RUN_ID / "output"
    finally:
        a._FILES_DIR = original_dir
        shared.last_architecture = ""
        shared.output_file = ""
        shared.output_file_url = ""


async def test_the_saved_docx_is_the_designed_document_with_only_the_requested_sections(design_writing):
    from docx import Document

    a, shared, llm, out_dir = design_writing

    await a.generate_architecture_from_context.ainvoke({"context": "ctx", "components": ["hld"]})

    files = sorted(out_dir.glob("*.docx"))
    assert files, "no .docx written"
    doc = Document(str(files[0]))
    texts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                texts.extend(p.text for p in cell.paragraphs)
    joined = "\n".join(texts)
    assert "DESIGN DOCUMENT" in joined, "the title band is the designed document's signature"
    assert "QuickLink" in joined
    assert "High-level design" in joined
    assert "Database schema" not in joined, "a section that was not produced must not appear"
    assert doc.core_properties.title == "QuickLink"


async def test_adding_a_component_rewrites_the_same_file_with_both_sections(design_writing):
    from docx import Document

    a, shared, llm, out_dir = design_writing
    await a.generate_architecture_from_context.ainvoke({"context": "ctx", "components": ["hld"]})
    llm.return_value = DB_DOC
    await a.generate_architecture_from_context.ainvoke({"context": "ctx", "components": ["db"]})

    files = sorted(out_dir.glob("*.docx"))
    assert len(files) == 1, [f.name for f in files]
    doc = Document(str(files[0]))
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "High-level design" in joined and "Database schema" in joined


# ── the menu tool and the prompt ─────────────────────────────────────────────


async def test_the_menu_tool_lists_the_components():
    from agents_orchestrator.design_architecture_agent.agents import architecture as a

    names = {t.name for t in a.tools}
    assert "list_design_components" in names
    out = await a.list_design_components.ainvoke({})
    assert "High-level design" in out and "Database schema" in out
    assert "full design document" in out.lower()


def test_the_system_prompt_no_longer_demands_all_eight():
    from agents_orchestrator.design_architecture_agent.agents import architecture as a
    from agents_orchestrator.design_architecture_agent import components as dc

    prompt = a.DESIGN_SYS_MESSAGE
    assert "ALL 8" not in prompt
    assert "all 8 section" not in prompt.lower()
    assert "list_design_components" in prompt
    for c in dc.COMPONENTS:
        if c.id != "overview":
            assert c.label in prompt, f"the prompt must name {c.label}"
    assert "only" in prompt.lower()
