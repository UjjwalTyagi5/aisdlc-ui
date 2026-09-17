"""The Design agent's Word document, in the brief's visual language.

The file used to be the generic markdown→docx output: default headings, "Table Grid"
tables, no title, no identity. `design_document.render_design_docx` paints the same
markdown with the palette and primitives the Track 3 brief uses (`shared/docs/pwc_style`).

What is pinned, by reading the produced .docx back with python-docx:

  * a title band (eyebrow, title, the components it holds, the project) and a
    key-facts strip precede the content;
  * only the sections the markdown holds appear, numbered, in order;
  * a markdown table's header row is shaded with the palette's white→orange steps
    and its body is not — one colour, one meaning;
  * fenced code (DDL, YAML, JSON) lands in a monospace panel, not as prose;
  * a Mermaid block becomes an image when the renderer returns PNG bytes, and falls
    back to the code, labelled, when it does not;
  * an `[ASSUMPTION]` blockquote becomes a callout the reader can spot;
  * bullets and bold/inline code survive.
"""
from __future__ import annotations

import io
import struct
import zlib
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from agents_orchestrator.design_architecture_agent.design_document import (
    DesignMeta,
    render_design_docx,
)
from shared.docs import pwc_style as st

SAMPLE = """\
# QuickLink Internal URL Shortener

## HIGH-LEVEL DESIGN

QuickLink is a server-rendered **Node.js + Express** application. The redirect path
is isolated for `p95 ≤ 150 ms`.

```mermaid
graph TD
    A["Client"] --> B["Express App"]
    B --> C[("SQLite")]
```

| System | Protocol | Auth | Purpose |
|--------|----------|------|---------|
| Intranet CMS | HTTPS | Gateway | Auto-create links |
| Corporate Network | HTTPS | Gateway | All traffic |
| link | id, store_id, price_cents, created_at | Gateway | _snake_case_ columns |

> ⚠️ **[ASSUMPTION]** — no cache was specified in requirements. Confirm before development.

- Redirect never blocks on click recording
- Disabled codes show a branded not-found page
- `Link` rows are never updated in place, only **appended**
- [ASSUMPTION] Peak load is under 200 clicks per minute

1. Slugs are generated server-side
2. Clicks are recorded asynchronously

## DATABASE SCHEMA

```sql
CREATE TABLE link (
    id INTEGER PRIMARY KEY,
    slug TEXT NOT NULL UNIQUE
);
```
"""


def _png() -> bytes:
    """A valid 1×1 PNG, so python-docx accepts it as a picture."""
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    idat = zlib.compress(b"\x00\xff\x00\x00")
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")


META = DesignMeta(
    title="QuickLink Internal URL Shortener",
    project="TEST Project",
    components=["hld", "db"],
    source="QuickLink_BRD_new.docx (approved)",
    generated_on="15 Sep 2026",
)


@pytest.fixture
def rendered(tmp_path):
    path = tmp_path / "design.docx"
    render_design_docx(SAMPLE, str(path), meta=META, render_mermaid=lambda code: _png())
    return Document(str(path))


def _texts(doc) -> list[str]:
    """Every paragraph once. A merged cell is reported by python-docx for each grid
    column it spans (the title band is 24 of them), so cells are deduplicated."""
    out = [p.text for p in doc.paragraphs]
    seen: list = []  # the elements themselves: keeping them referenced keeps identity stable
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if any(cell._tc is tc for tc in seen):
                    continue
                seen.append(cell._tc)
                out.extend(p.text for p in cell.paragraphs)
    return out


def _fills(table) -> list[list[str | None]]:
    fills = []
    for row in table.rows:
        row_fills = []
        for cell in row.cells:
            shd = cell._tc.tcPr.find(qn("w:shd")) if cell._tc.tcPr is not None else None
            row_fills.append(shd.get(qn("w:fill")) if shd is not None else None)
        fills.append(row_fills)
    return fills


def test_the_title_band_names_the_document_and_what_it_holds(rendered):
    text = "\n".join(_texts(rendered))
    assert "QuickLink Internal URL Shortener" in text
    assert "TEST Project" in text
    assert "High-level design" in text and "Database schema" in text
    assert "DESIGN" in text, "the eyebrow says what kind of document this is"


def test_the_key_facts_strip_carries_the_source_and_the_date(rendered):
    text = "\n".join(_texts(rendered))
    assert "QuickLink_BRD_new.docx (approved)" in text
    assert "15 Sep 2026" in text


def test_the_facts_strip_says_how_much_of_the_design_this_is(rendered):
    """A scoped document must read as scoped, not as unfinished: "2 of 8 components"
    is the fact a reviewer needs before wondering where the API contract went."""
    from agents_orchestrator.design_architecture_agent.design_document import scope_label

    text = "\n".join(_texts(rendered))
    assert "2 of 8 components" in text
    assert scope_label(["overview", "hld", "lld", "c4", "api", "db", "adr", "stack", "security"]) == "Full design document"
    assert scope_label(["hld", "lld", "c4", "api", "db", "adr", "stack", "security"]) == "Full design document"
    assert scope_label([]) == ""


def test_each_fact_appears_once(rendered):
    """The band and the facts strip split the facts between them; the generated
    date is a strip fact and the component list is the band's subtitle, and neither
    is said twice within three centimetres of itself."""
    text = "\n".join(_texts(rendered))
    assert text.count("15 Sep 2026") == 1
    assert text.count("High-level design, Database schema") == 1


def test_snake_case_identifiers_are_not_italicised(rendered):
    """`store_id, price_cents` is a column list. Underscore emphasis only applies at
    word edges, as in CommonMark — this rendered as "store<i>id, price</i>cents"."""
    runs = [r for p in rendered.paragraphs for r in p.runs]
    for t in rendered.tables:
        for row in t.rows:
            for cell in row.cells:
                runs.extend(r for p in cell.paragraphs for r in p.runs)
    joined = "".join(r.text for r in runs)
    assert "store_id, price_cents" in joined
    assert not any(r.italic and "id, price" in r.text for r in runs)
    assert any(r.italic and r.text == "snake_case" for r in runs), "real emphasis still works"


def test_an_inline_assumption_marker_becomes_a_tag(rendered):
    runs = [r for p in rendered.paragraphs for r in p.runs]
    tag = [r for r in runs if r.text.strip() == "ASSUMPTION" and r.bold and r.font.size.pt < 9]
    assert tag, "the [ASSUMPTION] marker paints as a small tag"
    assert "[ASSUMPTION]" not in "".join(r.text for r in runs)


def test_only_the_sections_present_appear_numbered_in_order(rendered):
    body = [p.text for p in rendered.paragraphs]
    heads = [t for t in body if t.startswith(("01", "02", "03"))]
    assert len(heads) == 2, heads
    assert "High-level design" in heads[0]
    assert "Database schema" in heads[1]
    assert not any("Low-level" in t for t in body)


def test_a_table_header_row_is_orange_tinted_and_the_body_is_not(rendered):
    table = next(t for t in rendered.tables if t.rows[0].cells[0].text.strip() == "System")
    fills = _fills(table)
    assert all(f in st.HEAD_STEPS for f in fills[0]), fills[0]
    assert all(f not in st.HEAD_STEPS for f in fills[1]), fills[1]
    assert table.rows[1].cells[0].text.strip() == "Intranet CMS"


def test_fenced_code_lands_in_a_monospace_panel(rendered):
    # The DDL is inside a shaded single-cell table, in the code face.
    panel = next(
        t for t in rendered.tables
        if "CREATE TABLE link" in t.rows[0].cells[0].text
    )
    fills = _fills(panel)
    assert fills[0][0] == st.PANEL
    runs = [r for p in panel.rows[0].cells[0].paragraphs for r in p.runs]
    assert runs and all(r.font.name == st.CODE_FONT for r in runs)


def test_a_mermaid_block_becomes_a_figure(rendered):
    assert len(rendered.inline_shapes) == 1
    text = "\n".join(_texts(rendered))
    assert "Figure 1" in text


def test_a_mermaid_block_falls_back_to_its_code_when_rendering_fails(tmp_path):
    path = tmp_path / "design.docx"
    render_design_docx(SAMPLE, str(path), meta=META, render_mermaid=lambda code: None)
    doc = Document(str(path))
    assert len(doc.inline_shapes) == 0
    text = "\n".join(_texts(doc))
    assert 'A["Client"] --> B["Express App"]' in text
    assert "could not be rendered" in text.lower()


def test_an_assumption_becomes_a_callout(rendered):
    text = "\n".join(_texts(rendered))
    assert "ASSUMPTION" in text
    assert "no cache was specified" in text
    # It sits in its own shaded cell with an amber edge — the brief's "needs
    # confirming" mark, never a fill of the brand colour.
    cell = next(
        t.rows[0].cells[0] for t in rendered.tables
        if "no cache was specified" in t.rows[0].cells[0].text
    )
    borders = cell._tc.tcPr.find(qn("w:tcBorders"))
    assert borders is not None
    left = borders.find(qn("w:left"))
    assert left is not None and left.get(qn("w:color")) == st.AMBER
    # The eyebrow already says ASSUMPTION; the body must not open with the marker
    # again, whether it was written as "⚠️ **[ASSUMPTION]** —" or plainly.
    assert cell.text.count("ASSUMPTION") == 1, cell.text
    assert "[ASSUMPTION]" not in cell.text


def test_a_plain_assumption_callout_loses_its_marker_too(tmp_path):
    path = tmp_path / "plain.docx"
    render_design_docx("# T\n\n## HIGH-LEVEL DESIGN\n\n> [ASSUMPTION] Loyalty is out of scope.\n",
                       str(path), meta=DesignMeta(title="T"))
    doc = Document(str(path))
    cell = next(t.rows[0].cells[0] for t in doc.tables if "Loyalty" in t.rows[0].cells[0].text)
    assert cell.text.count("ASSUMPTION") == 1, cell.text
    assert "Loyalty is out of scope." in cell.text


def test_bullets_and_inline_formatting_survive(rendered):
    text = "\n".join(_texts(rendered))
    assert "Redirect never blocks on click recording" in text
    bold = [r for p in rendered.paragraphs for r in p.runs if r.bold and "Node.js" in r.text]
    assert bold, "bold inline text must stay bold"


def test_inline_code_inside_a_bullet_keeps_its_own_size(rendered):
    """A bullet run may carry its own size (inline code is a point smaller): the
    bullet's default size must not clash with it — this was a TypeError that lost
    the whole document whenever a bullet mentioned a table or field in backticks."""
    code = [r for p in rendered.paragraphs for r in p.runs if r.text == "Link"]
    assert code and code[0].font.name == st.CODE_FONT
    assert code[0].font.size.pt == 9


def test_numbered_lists_are_numbered_and_bulleted_lists_are_not(rendered):
    text = "\n".join(_texts(rendered))
    assert "1.  Slugs are generated server-side" in text
    assert "2.  Clicks are recorded asynchronously" in text
    assert "\u2022  Redirect never blocks" in text


def test_the_document_has_page_numbers_and_metadata(rendered):
    footer_text = "".join(p.text for p in rendered.sections[0].footer.paragraphs)
    assert "Page" in footer_text
    assert rendered.core_properties.title == "QuickLink Internal URL Shortener"


def test_no_orange_text_uses_the_brand_fill_colour(rendered):
    """Brand orange is 3.3:1 on white — a fill or a rule, never text
    (accessibility: text needs 4.5:1). The renderer must never colour a run with it."""
    for p in rendered.paragraphs:
        for r in p.runs:
            if r.font.color is not None and r.font.color.rgb is not None:
                assert str(r.font.color.rgb) != st.BRAND
