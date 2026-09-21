"""Any document an agent filed opens on its page — a workbook as its sheets, a document as text.

LIVE (2026-09-21): opening a document from the Documents panel worked on four agent pages and
did nothing on five, and no page could show a spreadsheet — the Project Manager's estimate and
the Testing agent's workbooks were download-only. `GET /artifacts/{id}/preview` answers every
kind the same way, with the same guards as `/page`.
"""
from __future__ import annotations

import io
import sys
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from shared.services import sheet_preview  # noqa: E402
from shared.services.sheet_preview import can_preview_sheet, xlsx_sheets  # noqa: E402

TENANT = "11111111-1111-1111-1111-111111111111"
BASE = f"{TENANT}/bu/proj/plan/run/document"


def _xlsx(build) -> bytes:
    wb = Workbook()
    build(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _estimate(wb):
    ws = wb.active
    ws.title = "Estimate"
    ws.append(["Task", "Owner", "Days", "Start"])
    ws.append(["Link creation", "Dev", 5.0, datetime(2026, 10, 1)])
    ws.append(["Reporting", "Dev", 2.5, date(2026, 10, 8)])
    ws.append([None, None, None, None])                 # a trailing empty row
    notes = wb.create_sheet("Notes")
    notes.append(["Assumes one developer."])


# ── the workbook reader ───────────────────────────────────────────────────────


def test_a_workbook_reads_as_its_sheets_in_order_as_a_person_would_read_the_cells():
    sheets = xlsx_sheets(_xlsx(_estimate))
    assert [s["name"] for s in sheets] == ["Estimate", "Notes"]
    rows = sheets[0]["rows"]
    assert rows[0] == ["Task", "Owner", "Days", "Start"]
    # whole floats lose Excel's ".0"; a midnight datetime reads as a date
    assert rows[1] == ["Link creation", "Dev", "5", "2026-10-01"]
    assert rows[2][2] == "2.5"
    assert len(rows) == 3, "trailing empty rows are dropped"
    assert sheets[0]["truncated"] is False


def test_every_row_is_as_wide_as_the_widest_so_the_grid_has_straight_columns():
    def ragged(wb):
        ws = wb.active
        ws.append(["a"])
        ws.append(["b", "c", "d"])
    rows = xlsx_sheets(_xlsx(ragged))[0]["rows"]
    assert rows == [["a", "", ""], ["b", "c", "d"]]


def test_a_large_sheet_is_cut_and_says_so(monkeypatch):
    monkeypatch.setattr(sheet_preview, "MAX_ROWS", 5)

    def big(wb):
        for i in range(12):
            wb.active.append([f"row {i}"])
    sheet = xlsx_sheets(_xlsx(big))[0]
    assert len(sheet["rows"]) == 5 and sheet["rows_total"] == 12 and sheet["truncated"] is True


def test_bytes_that_are_not_a_workbook_have_no_view():
    assert xlsx_sheets(b"not a workbook") is None
    assert can_preview_sheet("QuickLink_Estimate.xlsx") and can_preview_sheet("x.XLSM")
    assert not can_preview_sheet("architecture.docx") and not can_preview_sheet("deck.pptx")


# ── the route ─────────────────────────────────────────────────────────────────


class _Store:
    def __init__(self, blobs: dict[str, bytes]):
        self.blobs = blobs

    async def download_bytes(self, name):
        if name not in self.blobs:
            raise FileNotFoundError(name)
        return self.blobs[name]


def _pending(path: str) -> str:
    from shared.services.artifact_store import pending_blob_path
    return pending_blob_path(path)


async def _preview(monkeypatch, filename: str, blobs: dict, status: str = "draft"):
    from shared.routers import artifacts as r

    art = SimpleNamespace(id="art-1", project_id="proj", blob_path=f"{BASE}/{filename}", approval_status=status)
    request = SimpleNamespace(state=SimpleNamespace(tenant_id=TENANT, user_id="u1", permissions=[]),
                              app=SimpleNamespace(state=SimpleNamespace(blob_client=_Store(blobs))))
    monkeypatch.setattr(r, "_get_artifact_or_404", AsyncMock(return_value=(art, None)))
    monkeypatch.setattr(r, "_assert_project_visible", AsyncMock())
    return await r.artifact_preview("art-1", request, db=None)


async def test_a_spreadsheet_opens_as_its_sheets_wherever_its_bytes_are(monkeypatch):
    data = _xlsx(_estimate)
    draft = await _preview(monkeypatch, "Estimate.xlsx", {_pending(f"{BASE}/Estimate.xlsx"): data})
    assert draft["kind"] == "sheets" and [s["name"] for s in draft["sheets"]] == ["Estimate", "Notes"]
    assert draft["filename"] == "Estimate.xlsx" and draft["status"] == "draft"

    approved = await _preview(monkeypatch, "Estimate.xlsx", {f"{BASE}/Estimate.xlsx": data}, status="approved")
    assert approved["kind"] == "sheets", "an approved workbook is read from its final path"


async def test_a_workbook_is_shown_as_itself_even_when_it_has_a_page_copy(monkeypatch):
    """The Testing agent's reports have a markdown page copy too; a person opening a
    spreadsheet expects the spreadsheet."""
    path = f"{BASE}/Unit_Test_Report.xlsx"
    out = await _preview(monkeypatch, "Unit_Test_Report.xlsx", {
        _pending(path): _xlsx(_estimate),
        _pending(f"{path}.page.md"): b"# Report",
    })
    assert out["kind"] == "sheets"


async def test_a_document_opens_as_its_page_copy_or_its_text(monkeypatch):
    from docx import Document

    path = f"{BASE}/BRD.docx"
    copy = await _preview(monkeypatch, "BRD.docx", {_pending(f"{path}.page.md"): b"## Scope\nShort links."})
    assert copy == {"artifactId": "art-1", "filename": "BRD.docx", "status": "draft",
                    "kind": "markdown", "markdown": "## Scope\nShort links.", "derived": False}

    doc, buffer = Document(), io.BytesIO()
    doc.add_heading("Architecture", level=1)
    doc.save(buffer)
    derived = await _preview(monkeypatch, "BRD.docx", {_pending(path): buffer.getvalue()})
    assert derived["kind"] == "markdown" and derived["derived"] is True and "# Architecture" in derived["markdown"]
    assert derived["derivedFrom"] == "word"


# ── every other kind an agent files: read from the file, only where there was no view ──


async def test_a_deck_opens_as_its_slides(monkeypatch):
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])          # title and content
    slide.shapes.title.text = "Architecture"
    body = slide.placeholders[1].text_frame
    body.text = "Three services"
    body.add_paragraph().text = "One database"
    table = deck.slides.add_slide(deck.slide_layouts[5])
    table.shapes.title.text = "Risks"
    grid = table.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    grid.cell(0, 0).text, grid.cell(0, 1).text = "Risk", "Level"
    grid.cell(1, 0).text, grid.cell(1, 1).text = "Latency", "Low"
    table.notes_slide.notes_text_frame.text = "Mention the SLA."
    buffer = io.BytesIO()
    deck.save(buffer)

    out = await _preview(monkeypatch, "Architecture.pptx", {_pending(f"{BASE}/Architecture.pptx"): buffer.getvalue()})
    assert out["kind"] == "markdown" and out["derived"] is True and out["derivedFrom"] == "slides"
    md = out["markdown"]
    assert "## Slide 1 · Architecture" in md and "- Three services\n- One database" in md
    assert "## Slide 2 · Risks" in md and "| Risk | Level |" in md and "| Latency | Low |" in md
    assert "> **Speaker notes:** Mention the SLA." in md


async def test_a_csv_a_markdown_file_and_a_text_file_open_as_themselves(monkeypatch):
    csv_out = await _preview(monkeypatch, "cases.csv", {_pending(f"{BASE}/cases.csv"): "id;title\r\nTC-1;Créer un lien\r\n".encode("utf-8-sig")})
    assert csv_out["kind"] == "sheets" and csv_out["sheets"][0]["name"] == "cases"
    assert csv_out["sheets"][0]["rows"] == [["id", "title"], ["TC-1", "Créer un lien"]], "delimiter sniffed, BOM dropped"

    md = await _preview(monkeypatch, "Effort.md", {_pending(f"{BASE}/Effort.md"): b"# Effort\n\n## Build\n5 days"})
    assert md["kind"] == "markdown" and md["markdown"].startswith("# Effort") and md["derived"] is False

    js = await _preview(monkeypatch, "config.json", {_pending(f"{BASE}/config.json"): b'{"a": "```"}'})
    assert js["kind"] == "markdown" and js["markdown"].startswith("````json\n"), "fenced past the text's own backticks"


async def test_an_html_report_is_handed_over_for_a_sandboxed_frame(monkeypatch):
    out = await _preview(monkeypatch, "coverage.html", {_pending(f"{BASE}/coverage.html"): b"<h1>Coverage 91%</h1>"})
    assert out == {"artifactId": "art-1", "filename": "coverage.html", "status": "draft",
                   "kind": "html", "html": "<h1>Coverage 91%</h1>", "truncated": False}


async def test_a_pdf_or_an_image_is_drawn_by_the_browser(monkeypatch):
    pdf = await _preview(monkeypatch, "BRD.pdf", {})
    assert pdf["kind"] == "file" and pdf["media"] == "pdf" and pdf["contentType"] == "application/pdf"
    image = await _preview(monkeypatch, "c4.svg", {})
    assert image["kind"] == "file" and image["media"] == "image"


async def test_a_page_copy_still_wins_over_the_file_so_nothing_that_opened_before_changes(monkeypatch):
    path = f"{BASE}/Design.pdf"
    out = await _preview(monkeypatch, "Design.pdf", {_pending(f"{path}.page.md"): b"## Overview\nBody."})
    assert out["kind"] == "markdown" and out["markdown"] == "## Overview\nBody." and out["derived"] is False


# ── the bytes behind a PDF or an image ─────────────────────────────────────────


async def _preview_file(monkeypatch, filename: str, blobs: dict, status: str = "draft"):
    from shared.routers import artifacts as r

    art = SimpleNamespace(id="art-1", project_id="proj", blob_path=f"{BASE}/{filename}", approval_status=status)
    request = SimpleNamespace(state=SimpleNamespace(tenant_id=TENANT, user_id="u1", permissions=[]),
                              app=SimpleNamespace(state=SimpleNamespace(blob_client=_Store(blobs))))
    monkeypatch.setattr(r, "_get_artifact_or_404", AsyncMock(return_value=(art, None)))
    monkeypatch.setattr(r, "_assert_project_visible", AsyncMock())
    return await r.artifact_preview_file("art-1", request, db=None)


async def test_a_pdf_is_served_inline_with_its_own_type_and_nosniff(monkeypatch):
    response = await _preview_file(monkeypatch, "BRD.pdf", {_pending(f"{BASE}/BRD.pdf"): b"%PDF-1.7"})
    assert response.body == b"%PDF-1.7" and response.media_type == "application/pdf"
    assert response.headers["content-disposition"] == 'inline; filename="BRD.pdf"'
    assert response.headers["x-content-type-options"] == "nosniff"
    assert "content-security-policy" not in response.headers


async def test_an_svg_cannot_run_script_even_opened_directly(monkeypatch):
    response = await _preview_file(monkeypatch, "c4.svg", {f"{BASE}/c4.svg": b"<svg/>"}, status="approved")
    assert response.media_type == "image/svg+xml"
    assert response.headers["content-security-policy"].startswith("sandbox;")


async def test_only_pdfs_and_images_are_served_this_way(monkeypatch):
    with pytest.raises(HTTPException) as word:
        await _preview_file(monkeypatch, "BRD.docx", {_pending(f"{BASE}/BRD.docx"): b"PK"})
    assert word.value.status_code == 404
    with pytest.raises(HTTPException) as rejected:
        await _preview_file(monkeypatch, "BRD.pdf", {}, status="rejected")
    assert rejected.value.status_code == 410
    with pytest.raises(HTTPException) as lost:
        await _preview_file(monkeypatch, "BRD.pdf", {})
    assert lost.value.status_code == 404 and "could not be found in storage" in lost.value.detail


async def test_a_file_with_no_view_a_rejected_one_and_a_missing_one_say_so(monkeypatch):
    with pytest.raises(HTTPException) as archive:
        await _preview(monkeypatch, "bundle.zip", {_pending(f"{BASE}/bundle.zip"): b"PK"})
    assert archive.value.status_code == 404 and "download" in archive.value.detail

    with pytest.raises(HTTPException) as unreadable:
        await _preview(monkeypatch, "Architecture.pptx", {_pending(f"{BASE}/Architecture.pptx"): b"not a deck"})
    assert unreadable.value.status_code == 404 and "could not be read" in unreadable.value.detail

    with pytest.raises(HTTPException) as lost:
        await _preview(monkeypatch, "Brief.docx", {})
    assert lost.value.status_code == 404 and "could not be found in storage" in lost.value.detail

    with pytest.raises(HTTPException) as rejected:
        await _preview(monkeypatch, "Estimate.xlsx", {}, status="rejected")
    assert rejected.value.status_code == 410

    with pytest.raises(HTTPException) as gone:
        await _preview(monkeypatch, "Estimate.xlsx", {})
    assert gone.value.status_code == 404 and "could not be read" in gone.value.detail


def test_the_preview_route_has_the_same_permission_as_the_page_route():
    from shared.routers import artifacts as r

    routes = {route.path: route for route in r.artifacts_router.routes}
    page = routes["/artifacts/{artifact_id}/page"]
    for path in ("/artifacts/{artifact_id}/preview", "/artifacts/{artifact_id}/preview/file"):
        assert [d.call.__name__ for d in routes[path].dependant.dependencies] == \
            [d.call.__name__ for d in page.dependant.dependencies], path
