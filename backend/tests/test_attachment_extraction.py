"""An attachment that cannot be read must SAY so, not send the agent chasing a path.

From a live Requirements chat. The user attached
"Screenshot 2026-09-02 004837.png" and the agent replied:

    Error: local file not found

The upload had worked perfectly. What failed is subtler and worse than a missing file:
`extract_file_text` returns a readable PLACEHOLDER when it cannot parse something —
"[Binary file: shot.png]" — and the caller tested it with `if _txt and _txt.strip()`.
A placeholder is a non-empty string, so the screenshot was announced to the agent as
"the user attached this, use its content directly" with the content being the words
"[Binary file: shot.png]". The agent then did the reasonable thing and tried to open
the file itself, and `upload_file` answered "local file not found" about a path it had
guessed out of prose.

So the user saw a file-not-found error for a file that uploaded fine, about a file type
the agent could never have read anyway. Every part of that message was misleading.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


# ── the sentinel ─────────────────────────────────────────────────────────────


@pytest.mark.unit
@pytest.mark.parametrize(
    "placeholder",
    [
        "[Binary file: Screenshot 2026-09-02 004837.png]",
        "[Binary file: diagram.svg]",
        "[Error reading notes.pdf: PdfReadError]",
    ],
)
def test_a_placeholder_is_not_content(placeholder):
    """The whole bug in one assertion: these are all truthy."""
    from shared.tools.document_tools import extraction_succeeded

    assert bool(placeholder.strip()) is True   # why `if text:` passed
    assert extraction_succeeded(placeholder) is False


@pytest.mark.unit
@pytest.mark.parametrize("text", ["", "   ", "\n\n", None])
def test_empty_extraction_is_not_content_either(text):
    from shared.tools.document_tools import extraction_succeeded

    assert extraction_succeeded(text) is False


@pytest.mark.unit
def test_real_text_is_content():
    from shared.tools.document_tools import extraction_succeeded

    assert extraction_succeeded("As a user, I want to log in so that ...") is True


@pytest.mark.unit
def test_text_merely_mentioning_a_binary_file_is_still_content():
    """The check is a PREFIX on the whole string, not a substring search — a document
    that happens to discuss binary files must not be discarded as unreadable."""
    from shared.tools.document_tools import extraction_succeeded

    assert extraction_succeeded("The build emits a [Binary file: app.exe] artifact.") is True


# ── extraction really does return a placeholder for an image ─────────────────


@pytest.mark.unit
def test_a_png_extracts_to_a_placeholder_not_to_text(tmp_path):
    """Pins the actual behaviour the callers now depend on, rather than assuming it."""
    from shared.tools.document_tools import extract_file_text, extraction_succeeded

    png = tmp_path / "Screenshot 2026-09-02 004837.png"
    png.write_bytes(bytes([0x89]) + b"PNG" + bytes([13, 10, 26, 10]) + bytes(32))

    out = extract_file_text(str(png))
    assert out.strip() != ""              # non-empty, which is what fooled the caller
    assert extraction_succeeded(out) is False


@pytest.mark.unit
def test_an_unreadable_attachment_names_the_limit_instead_of_a_path(tmp_path):
    """The old hint was "please use the following files <path>", which pointed the agent
    at a file tool that cannot read an image either. The replacement has to (a) stop the
    agent opening it, (b) stop it pretending it looked, and (c) say what WOULD work."""
    from shared.tools.document_tools import attachment_message_contents

    png = tmp_path / "diagram.png"
    png.write_bytes(bytes([0x89]) + b"PNG" + bytes([13, 10, 26, 10]) + bytes(32))

    message = attachment_message_contents([str(png)])[0]

    assert "could not be read as text" in message
    assert "do not call a file tool on it" in message
    assert "do not claim to have looked at it" in message
    for fmt in (".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx"):
        assert fmt in message, f"{fmt} not offered as an alternative"
    assert "please use the following files" not in message


# -- presentations ------------------------------------------------------------


def test_a_pptx_is_accepted_for_upload():
    """IT WAS NOT, and `chat_artifacts` generates decks in that very format — so the
    platform produced a file its own upload refused, and a design deck, the most obvious
    thing to attach to a Design agent, was the one thing you could not attach."""
    from shared.services.attachment_store import ALLOWED_ATTACHMENT_EXTS

    assert ".pptx" in ALLOWED_ATTACHMENT_EXTS


def test_ppt_is_still_refused():
    """DELIBERATE, not an oversight. python-pptx cannot read the legacy binary format,
    so accepting `.ppt` would store a file no agent could ever read — worse than
    refusing it, because the refusal at least says why."""
    from shared.services.attachment_store import ALLOWED_ATTACHMENT_EXTS

    assert ".ppt" not in ALLOWED_ATTACHMENT_EXTS


def test_a_pptx_extracts_slide_by_slide(tmp_path):
    """ACCEPTING IT IS HALF THE JOB. A format that uploads but extracts to
    "[Binary file: deck.pptx]" is the exact failure this file was written about: the
    placeholder is a non-empty string, so the agent is told to use it as content.
    """
    from pptx import Presentation

    from shared.tools.document_tools import extract_file_text, extraction_succeeded

    deck = Presentation()
    for title, body in (("Delivery Risks", "Vendor lock-in"), ("Timeline", "Q3 pilot")):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    path = tmp_path / "deck.pptx"
    deck.save(str(path))

    text = extract_file_text(str(path))

    assert extraction_succeeded(text), text
    assert "Delivery Risks" in text and "Q3 pilot" in text
    # ORDERING IS PART OF THE MEANING — "the risks slide" is a thing people say, so the
    # slide boundaries survive rather than being flattened into one blob.
    assert "--- Slide 1 ---" in text and "--- Slide 2 ---" in text
    assert text.index("Delivery Risks") < text.index("Timeline")
