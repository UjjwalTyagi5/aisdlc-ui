"""The migration-intent brief and the Track 3 Requirements agent's tools."""
from __future__ import annotations

import pytest

from agents_orchestrator.requirements_modernization_agent.brief import (
    REQUIRED_SECTIONS,
    brief_markdown,
    missing_sections,
)
from agents_orchestrator.requirements_modernization_agent.tools import brief_tools
from config.ws_helper import set_session_id, set_tenant_id, set_user_id
from shared.models.artifacts import MigrationIntentArtifact

COMPLETE = {
    "system_name": "Billing",
    "business_drivers": [".NET Framework 4.5.2 is out of support", "IIS hosting cost"],
    "current_stack": ".NET Framework 4.5.2 WebForms on IIS",
    "target_stack": ".NET 8 on Azure App Service",
    "in_scope": ["Billing.Web", "Billing.Core"],
    "out_of_scope": ["Reporting database"],
    "constraints": ["Live before 31 March", "Invoice PDF layout must not change"],
    "success_criteria": ["Identical invoices for the recorded Q1 inputs", "p95 under 300 ms"],
    "legacy_repository_url": "https://dev.azure.com/acme/billing/_git/billing",
}


@pytest.fixture(autouse=True)
def _ctx(tmp_path):
    set_session_id(f"req-{tmp_path.name}")
    set_user_id("")
    set_tenant_id(None)
    yield
    brief_tools._LAST_BRIEF.clear()


def test_an_empty_brief_is_missing_every_required_section():
    assert missing_sections(MigrationIntentArtifact()) == [label for _, label in REQUIRED_SECTIONS]


def test_whitespace_is_not_an_answer():
    brief = MigrationIntentArtifact(system_name="  ", business_drivers=["  "])
    assert "System being modernized" in missing_sections(brief)
    assert "Why this modernization is happening" in missing_sections(brief)


def test_markdown_is_a_document_with_the_intake_sections():
    brief = MigrationIntentArtifact(
        system_name="Billing", business_drivers=["EOL"],
        current_state={"stack": ".NET Framework 4.5.2"}, target_state={"stack": ".NET 8"},
        in_scope=["Billing.Web"], constraints=["March"], success_criteria=["Same invoices"],
    )
    md = brief_markdown(brief)
    assert md.startswith("# Migration Intent Brief — Billing")
    for heading in ("The change at a glance", "Why we are modernizing", "Scope", "Constraints",
                    "How we will measure success", "## Legacy repository"):
        assert heading in md
    # A version-1 brief (no per-layer change) shows today -> target as one row.
    assert "| **Whole system** | .NET Framework 4.5.2 | .NET 8 | — |" in md


async def test_an_incomplete_brief_is_not_persisted(monkeypatch):
    calls: list = []

    async def fake_persist(brief):
        calls.append(brief)
        return "saved"

    monkeypatch.setattr(brief_tools, "_persist", fake_persist)
    partial = {k: v for k, v in COMPLETE.items() if k != "success_criteria"}
    out = await brief_tools.record_migration_intent.ainvoke(partial)
    assert out.startswith("NOT RECORDED YET")
    assert "Success criteria" in out
    assert calls == []


async def test_a_complete_brief_is_persisted_and_returned_as_the_document(monkeypatch):
    calls: list = []

    async def fake_persist(brief):
        calls.append(brief)
        return "Saved to the project as the current migration-intent brief."

    monkeypatch.setattr(brief_tools, "_persist", fake_persist)
    out = await brief_tools.record_migration_intent.ainvoke(COMPLETE)
    assert out.startswith("# Migration Intent Brief — Billing")
    assert len(calls) == 1 and calls[0].target_state.stack == ".NET 8 on Azure App Service"
    assert calls[0].legacy_repository.url.endswith("/_git/billing")
    assert calls[0].recorded_at


async def test_export_needs_a_recorded_brief():
    out = await brief_tools.export_migration_brief.ainvoke({})
    assert "record it first" in out


async def test_a_board_write_with_no_actor_is_refused_before_any_connector_call(monkeypatch):
    class _Board:
        display_name = "Azure DevOps"
        access_level = "read_write"

        async def write_adapter(self, *a, **k):
            raise AssertionError("must not write without an approving owner")

    monkeypatch.setattr("config.connectors.context.get_connector", lambda: _Board())
    out = await brief_tools.create_migration_work_items.ainvoke(
        {"project": "Billing", "epic_title": "Modernize Billing", "items_json": "[]"}
    )
    assert not out.startswith("Created")
    assert "needs a signed-in approver" in out


async def test_a_board_write_on_a_read_only_stage_names_the_grant(monkeypatch):
    class _ReadOnly:
        display_name = "Azure DevOps"
        access_level = "read"

    monkeypatch.setattr("config.connectors.context.get_connector", lambda: _ReadOnly())
    out = await brief_tools.create_migration_work_items.ainvoke(
        {"project": "Billing", "epic_title": "Modernize Billing"}
    )
    assert "cannot be used to write" in out


def test_prompt_is_migration_intent_not_stories():
    from agents_orchestrator.requirements_modernization_agent.prompts.migration_intent_prompt import (
        MIGRATION_INTENT_SYS_MESSAGE,
    )

    text = MIGRATION_INTENT_SYS_MESSAGE
    assert "MIGRATION-INTENT" in text
    assert "do NOT write user stories" in text
    assert "Success criteria" in text
    assert "Never invent" in text


def test_prompt_keeps_the_agents_own_examples_out_of_the_brief():
    """Found live (2026-09-10): asked for scope, the agent listed examples (data access
    layer, background jobs, reporting); the user named three items and three exclusions;
    the recorded brief carried the user's items PLUS the agent's examples, with "(if any)"
    and "unless ..." attached. The rule has to name that exact failure."""
    from agents_orchestrator.requirements_modernization_agent.prompts.migration_intent_prompt import (
        MIGRATION_INTENT_SYS_MESSAGE,
    )

    text = MIGRATION_INTENT_SYS_MESSAGE
    assert "Record ONLY what the user said or explicitly confirmed" in text
    assert "prompts, not answers" in text
    assert "(if any)" in text and "open questions" in text


def test_graph_compiles():
    from agents_orchestrator.requirements_modernization_agent.agents.intake import TOOLS, app

    assert app is not None
    assert {t.name for t in TOOLS} == {
        "record_migration_intent", "export_migration_brief", "list_board_projects",
        "create_migration_work_items",
        "get_legacy_code_profile", "list_legacy_files", "read_legacy_file", "search_legacy_code",
        "find_legacy_repositories", "pull_legacy_code",
    }


# ── version 2: the structured brief ─────────────────────────────────────────

V2 = {
    "system_name": "ClaimTrack",
    "goal": "Move ClaimTrack onto supported runtimes by June 2027.",
    "drivers": [{"category": "End of life", "title": "Java 7 is unsupported", "detail": "Since 2022."},
                {"category": "cost", "title": "Lease ends", "detail": "30 June 2027."}],
    "layers": [
        {"layer": "Settlement batch", "current": "Java 7", "target": "Java 21 + Spring Batch",
         "change_type": "Upgrade", "modules": ["claimtrack-batch"]},
        {"layer": "Broker portal", "current": "AngularJS 1.5", "target": "React 18 + TypeScript",
         "change_type": "rewrite", "modules": ["agent-portal"]},
    ],
    "recommendation_summary": "Upgrade Java in place; rebuild only the portal.",
    "recommendation_rationale": ["AngularJS has no upgrade path."],
    "alternatives": [{"option": "Lift and shift", "why_not": "Keeps Java 7."}],
    "module_changes": [
        {"module": "claimtrack-batch", "target": "Java 21", "change_type": "upgrade", "effort": "Medium",
         "changes": ["Quartz -> Spring Batch"]},
        {"module": "agent-portal", "target": "React 18", "change_type": "rewrite", "effort": "high",
         "changes": ["Rebuild the four screens"]},
    ],
    "trade_offs": [{"decision": "Rebuild the portal", "gain": "Hireable", "cost": "Largest piece of work"}],
    "in_scope": ["claimtrack-batch", "agent-portal"],
    "constraints": ["Off the old servers by 30 June 2027"],
    "deadline": "2027-06-30", "budget": "$450,000",
    "milestones": [{"date": "2027-06-30", "label": "Data-centre exit", "kind": "Data-centre exit"},
                   {"date": "2027-02-01", "label": "Legacy freeze", "kind": "freeze"}],
    "success_measures": [{"metric": "API p95", "current": "~800 ms", "target": "<= 300 ms"}],
    "success_criteria": ["Identical payouts on 10,000 recorded claims"],
}


def test_labels_are_normalised_so_the_views_can_colour_them():
    brief = MigrationIntentArtifact(
        drivers=[{"category": "SOC 2 compliance deadline", "title": "t"}],
        layers=[{"layer": "x", "current_status": "End of life", "change_type": "Re-platform"}],
        module_changes=[{"module": "m", "effort": "Large", "change_type": "Lift and shift"}],
        milestones=[{"date": "2027-06-30", "label": "exit", "kind": "Data-centre exit"}],
        recommendation={"recommended_by": "BA"},
    )
    assert brief.drivers[0].category == "compliance"
    assert (brief.layers[0].current_status, brief.layers[0].change_type) == ("eol", "replatform")
    assert (brief.module_changes[0].effort, brief.module_changes[0].change_type) == ("high", "replatform")
    assert brief.milestones[0].kind == "deadline"
    assert brief.recommendation.recommended_by == "user"


def test_a_version_1_brief_still_loads():
    old = MigrationIntentArtifact(**{"system_name": "Billing", "business_drivers": ["EOL"], "version": 1})
    assert old.version == 1 and old.layers == [] and old.recommendation is None


async def test_a_structured_brief_is_recorded_with_the_plain_fields_derived(monkeypatch):
    stored: list = []

    async def fake_persist(brief):
        stored.append(brief)
        return "Saved to the project as the current migration-intent brief."

    monkeypatch.setattr(brief_tools, "_persist", fake_persist)
    out = await brief_tools.record_migration_intent.ainvoke(V2)
    assert not out.startswith("NOT RECORDED"), out
    brief = stored[0]
    # The required plain fields come from the structured ones.
    assert brief.business_drivers == ["Java 7 is unsupported: Since 2022.", "Lease ends: 30 June 2027."]
    assert brief.current_state.stack == "Settlement batch: Java 7; Broker portal: AngularJS 1.5"
    assert brief.target_state.stack.startswith("Settlement batch: Java 21")
    assert brief.recommendation.recommended_by == "agent"
    assert brief.drivers[0].category == "end_of_support"
    assert brief.module_changes[0].effort == "medium"
    # The document the tool returns is the numbered brief.
    for heading in ("## 1. The change at a glance", "Recommended target stack", "What changes in each module",
                    "Trade-offs", "Timeline", "How we will measure success"):
        assert heading in out
    assert "| **Settlement batch** | Java 7 | Java 21 + Spring Batch | Upgrade |" in out
    timeline = out[out.index("Timeline"):]
    assert timeline.index("1 Feb 2027") < timeline.index("30 Jun 2027")  # milestones in date order


async def test_todays_facts_come_from_the_pulled_code(monkeypatch):
    """The model says which modules change; the pulled code says what they run on today
    and whether it is still supported — the model cannot get that wrong."""
    from agents_orchestrator.modernization_common import legacy_code

    profile = {"modules": [
        {"name": "claimtrack-batch", "path": "claimtrack-batch", "runtime": "Java 7", "runtimeStatus": "eol"},
        {"name": "claimtrack-agent-portal", "path": "agent-portal", "runtime": "Node.js 8", "runtimeStatus": "eol"},
    ]}
    monkeypatch.setattr(legacy_code, "current_scope", lambda: ("proj", None))
    monkeypatch.setattr(legacy_code, "current_pull",
                        lambda project_id, run_id=None: {"profile": profile, "url": "https://x/y", "name": "y"})
    stored: list = []

    async def fake_persist(brief):
        stored.append(brief)
        return "saved"

    monkeypatch.setattr(brief_tools, "_persist", fake_persist)
    await brief_tools.record_migration_intent.ainvoke({
        **V2, "module_changes": V2["module_changes"] + [
            {"module": "database", "target": "MySQL 8.0", "change_type": "replatform", "changes": ["x"]}],
    })
    brief = stored[0]
    # Only the code's modules are "modules"; the database change lives in the change table.
    assert [m.module for m in brief.module_changes] == ["claimtrack-batch", "agent-portal"]
    portal = next(m for m in brief.module_changes if m.module == "agent-portal")
    assert (portal.current, portal.current_status, portal.path) == ("Node.js 8", "eol", "agent-portal")
    assert [layer.current_status for layer in brief.layers] == ["eol", "eol"]


def _v2_brief():
    from agents_orchestrator.requirements_modernization_agent.tools.brief_tools import MigrationIntentArtifact as M

    return M(system_name="ClaimTrack", goal=V2["goal"], drivers=V2["drivers"], business_drivers=["x"],
             current_state={"stack": "Java 7"}, target_state={"stack": "Java 21"}, layers=V2["layers"],
             recommendation={"summary": V2["recommendation_summary"], "rationale": V2["recommendation_rationale"],
                             "alternatives": V2["alternatives"]},
             module_changes=V2["module_changes"], trade_offs=V2["trade_offs"], in_scope=V2["in_scope"],
             constraints=V2["constraints"], deadline=V2["deadline"], budget=V2["budget"],
             milestones=V2["milestones"], success_measures=V2["success_measures"],
             success_criteria=V2["success_criteria"], stakeholders=[{"name": "Priya Raman", "role": "Owner"}])


def test_the_word_document_is_the_designed_brief(tmp_path):
    from docx import Document

    from agents_orchestrator.requirements_modernization_agent.brief_document import render_brief

    path = render_brief(_v2_brief(), str(tmp_path / "brief.docx"), meta={"version": 3, "status": "draft"})
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "ClaimTrack" in cells and "Version 3" in cells  # the title band is a table cell
    for title in ("The change at a glance", "Why we are modernizing", "Recommended target stack",
                  "What changes in each module", "Trade-offs", "Timeline", "How we will measure success"):
        assert title in text
    assert "Java 21 + Spring Batch" in cells and "End of life" not in cells  # no status: none given
    assert "RECOMMENDED BY THE REQUIREMENTS AGENT" in cells
    assert len(doc.inline_shapes) == 1  # the timeline chart


def test_the_pdf_is_the_designed_brief(tmp_path):
    from pypdf import PdfReader

    from agents_orchestrator.requirements_modernization_agent.brief_document import render_brief

    path = render_brief(_v2_brief(), str(tmp_path / "brief.pdf"))
    text = "\n".join(page.extract_text() for page in PdfReader(path).pages)
    for needle in ("ClaimTrack", "The change at a glance", "Recommended target stack", "Trade-offs",
                   "Data-centre exit", "How we will measure success"):
        assert needle in text, needle


def test_a_version_1_brief_renders_in_both_formats(tmp_path):
    from agents_orchestrator.requirements_modernization_agent.brief_document import render_brief

    old = MigrationIntentArtifact(system_name="Billing", business_drivers=["EOL"],
                                  current_state={"stack": ".NET 4.5"}, target_state={"stack": ".NET 8"},
                                  in_scope=["web"], constraints=["March"], success_criteria=["same"])
    for ext in ("docx", "pdf"):
        assert (tmp_path / f"b.{ext}").exists() is False
        render_brief(old, str(tmp_path / f"b.{ext}"))
        assert (tmp_path / f"b.{ext}").stat().st_size > 5_000


async def test_the_page_download_uses_the_designed_brief(monkeypatch, tmp_path):
    """`GET .../migration-intent/versions/{v}/export` renders THAT version with the brief
    renderer (title band, change table) — not the generic Markdown one."""
    from types import SimpleNamespace

    from docx import Document

    import shared.routers.modernization as router
    from shared.services import artifact_versions as svc

    async def guard(db, request, project_id, stage):
        return "proj", "tenant", "user"

    async def get_version(db, project_id, stage, version):
        return SimpleNamespace(payload=_v2_brief().model_dump(), status="published")

    monkeypatch.setattr(router, "_guard", guard)
    monkeypatch.setattr(svc, "get_version", get_version)
    resp = await router.export_version("proj", "migration-intent", 4, None, format="docx", db=None)
    assert resp.filename == "migration-intent-brief-v4.docx"
    cells = "\n".join(c.text for t in Document(resp.path).tables for r in t.rows for c in r.cells)
    assert "Version 4" in cells and "Approved" in cells
