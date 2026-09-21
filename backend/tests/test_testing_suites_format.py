"""The suite workbook is both a document and the input to a run: what is written is read back."""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from agents_orchestrator.testing_agent.suites.excel import SuiteFormatError, read_suite, write_suite  # noqa: E402
from agents_orchestrator.testing_agent.suites.models import SuiteMeta, validate_cases  # noqa: E402
from agents_orchestrator.testing_agent.suites.page import suite_markdown  # noqa: E402

META = dict(project="Url Shortner 1", source_project="QUICKLINK(Url shortner)", repository="QUICKLINK(Url shortner)", branch="feature/x", commit="082f91e49bf0",
            generated_at="17 Sep 2026 12:00 UTC", sources="Url_Shortner_1_BRD_v2.docx; architecture.docx; the code")

UNIT = [
    {"id": "UT-001", "title": "Creates a short link", "module": "src/services/urlService.js", "function": "createShortLink",
     "scenario": "happy", "input": "https://intranet/x", "expected": "Returns a 6-8 character code", "priority": "p1",
     "requirement": "FR-1"},
    {"id": "UT-001", "title": "Rejects an empty URL", "module": "src/services/urlService.js", "function": "createShortLink",
     "scenario": "negative", "input": "", "expected": "Throws a validation error"},
    {"id": "UT-003", "title": "Nothing to test", "module": "", "expected": ""},
]
FUNCTIONAL = [
    {"id": "FT-001", "title": "Create a short link from the home page", "requirement": "FR-1", "preconditions": "App running",
     "steps": [{"action": "navigate", "target": "/"}, {"action": "fill", "target": "Long URL", "value": "https://intranet/x"},
               {"action": "click", "target": "Create short link"}, {"action": "see", "value": "Short link created"}],
     "expected": "The new short link is shown"},
    {"id": "FT-002", "title": "Broken step", "steps": [{"action": "click"}], "expected": "x"},
]
API = [
    {"id": "AT-001", "title": "Create a link", "method": "post ", "path": "api/links", "headers": {"Content-Type": "application/json"},
     "body": {"url": "https://intranet/x"}, "expected_status": "201 Created", "expected_body_contains": "code; url",
     "capture": {"code": "code"}},
    {"id": "AT-002", "title": "Stats for the link", "method": "GET", "path": "/api/links/{{code}}/stats",
     "expected_status": 200, "expected_body_contains": ["clicks"]},
    {"id": "AT-003", "title": "No path", "method": "GET", "path": "", "expected_status": 200},
]


def test_invalid_cases_are_reported_and_ids_are_unique():
    cases, problems = validate_cases("unit", UNIT)
    assert [c.id for c in cases] == ["UT-001", "UT-001-2"]
    assert cases[0].scenario == "Happy path" and cases[1].scenario == "Error" and cases[0].priority == "High"
    assert len(problems) == 1 and problems[0].startswith("UT-003")


@pytest.mark.parametrize("kind,raw,valid", [("unit", UNIT, 2), ("functional", FUNCTIONAL, 1), ("api", API, 2)])
def test_a_suite_round_trips_through_its_workbook(tmp_path, kind, raw, valid):
    cases, _ = validate_cases(kind, raw)
    assert len(cases) == valid
    meta = SuiteMeta(kind=kind, **META)
    data = write_suite(meta, cases, str(tmp_path / "suite.xlsx"))
    meta2, cases2, problems = read_suite(data)
    assert problems == [] and meta2 == meta
    assert [c.model_dump() for c in cases2] == [c.model_dump() for c in cases]


def test_functional_steps_are_normalised_and_run_from_the_steps_sheet():
    cases, _ = validate_cases("functional", FUNCTIONAL)
    steps = cases[0].steps
    assert [s.action for s in steps] == ["open", "type", "click", "assert_text"]
    wb = load_workbook(io.BytesIO(write_suite(SuiteMeta(kind="functional", **META), cases)))
    assert wb.sheetnames == ["About", "Test cases", "Steps"]
    assert wb["Test cases"]["E2"].value.startswith('1. Open /\n2. Type "https://intranet/x" into "Long URL"')


def test_an_edited_workbook_runs_what_the_tester_changed():
    cases, _ = validate_cases("api", API)
    wb = load_workbook(io.BytesIO(write_suite(SuiteMeta(kind="api", **META), cases)))
    ws = wb["Test cases"]
    headers = [c.value for c in ws[1]]
    ws.cell(row=2, column=headers.index("Expected status") + 1, value="200")
    # Columns reordered by the tester: found by header, not position.
    ws.move_range("A1:A3", rows=0, cols=20)
    buf = io.BytesIO()
    wb.save(buf)
    _meta, edited, _ = read_suite(buf.getvalue())
    assert edited[0].expected_status == 200 and edited[0].capture == {"code": "code"}


def test_a_workbook_that_is_not_a_suite_is_refused_with_the_reason():
    from openpyxl import Workbook

    wb = Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(SuiteFormatError, match="not a test case suite"):
        read_suite(buf.getvalue())
    with pytest.raises(SuiteFormatError, match="not a readable Excel"):
        read_suite(b"not excel")


def test_the_page_view_is_the_suite():
    cases, _ = validate_cases("functional", FUNCTIONAL)
    md = suite_markdown(SuiteMeta(kind="functional", **META), cases)
    assert md.startswith("# QUICKLINK(Url shortner) — Functional test cases")
    assert "### FT-001 · Create a short link from the home page" in md and '4. Check the page shows "Short link created"' in md


def test_a_functional_case_that_checks_nothing_is_rejected():
    """LIVE: nine generated cases opened, typed and clicked, and verified nothing."""
    cases, problems = validate_cases("functional", [
        {"id": "FT-001", "title": "No check", "steps": [{"action": "open", "target": "/"}, {"action": "click", "target": "Go"}],
         "expected": "Link created"},
    ])
    assert cases == [] and "must end with an assert_text or assert_url step" in problems[0]


def test_an_api_case_may_only_use_a_variable_an_earlier_case_captured():
    cases, problems = validate_cases("api", [
        {"id": "AT-001", "title": "Stats before create", "method": "GET", "path": "/api/links/{{code}}/stats", "expected_status": 200},
        {"id": "AT-002", "title": "Create", "method": "POST", "path": "/api/shorten", "expected_status": 201, "capture": {"code": "shortCode"}},
        {"id": "AT-003", "title": "Stats", "method": "GET", "path": "/api/links/{{ code }}/stats", "expected_status": 200},
    ])
    assert [c.id for c in cases] == ["AT-002", "AT-003"]
    assert problems == ["AT-001: uses {{code}} before any earlier case captures it"]
