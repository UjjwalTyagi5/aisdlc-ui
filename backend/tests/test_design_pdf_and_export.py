"""The Design agent's PDF is the designed Word document, converted — with a fallback.

The Word file carries the platform's document identity (title band, facts, palette);
a PDF rendered from markdown by a second, plain renderer would not match it. So the PDF
is the .docx converted by Word (`docx2pdf`), and when Word is not there — a Linux host,
a locked COM session — the plain renderer still produces a readable PDF rather than an
error. `export_document(...docx)` goes through the same designed renderer for the same
reason.
"""
from __future__ import annotations

from unittest.mock import AsyncMock, MagicMock, patch

import pytest

RUN_ID = "11111111-2222-3333-4444-555555555555"
USER_ID = "af57932d-f4f2-4673-aef7-d33b75c022f4"
DOC = "# Coffee Ordering App\n\n## HIGH-LEVEL DESIGN\n\nThe storefront talks to an order service.\n"


@pytest.fixture
def design(tmp_path):
    from agents_orchestrator.design_architecture_agent.agents import architecture as a
    from agents_orchestrator.design_architecture_agent.config import shared
    from shared.services import chat_artifacts

    original_dir = a._FILES_DIR
    a._FILES_DIR = str(tmp_path)
    shared.last_architecture = DOC
    fake_manager = MagicMock()
    fake_manager.broadcast = AsyncMock()
    try:
        with patch.object(a, "get_user_id", lambda: USER_ID), \
                patch.object(a, "get_session_id", lambda: RUN_ID), \
                patch.object(a, "manager", fake_manager), \
                patch.object(a, "_render_mermaid_to_png", lambda code: None), \
                patch.object(chat_artifacts, "register_generated_file", AsyncMock()):
            yield a, tmp_path / USER_ID / "orchestrator" / RUN_ID / "output"
    finally:
        a._FILES_DIR = original_dir
        shared.last_architecture = ""


def _fake_convert(docx_path: str, pdf_path: str) -> None:
    from pathlib import Path

    assert Path(docx_path).exists(), "the designed .docx must exist before conversion"
    Path(pdf_path).write_bytes(b"%PDF-1.4 fake from " + Path(docx_path).name.encode())


async def test_the_pdf_is_converted_from_the_designed_docx(design):
    a, out_dir = design
    with patch.object(a, "_docx_to_pdf", _fake_convert):
        out = await a.save_architecture_pdf.ainvoke({"filename": "coffee.pdf"})

    pdf = out_dir / "coffee.pdf"
    assert pdf.exists() and pdf.read_bytes().startswith(b"%PDF")
    assert b"coffee.pdf-source.docx" in pdf.read_bytes(), "converted from the designed .docx"
    assert sorted(f.name for f in out_dir.iterdir()) == ["coffee.pdf"], (
        "the scratch Word file is removed: nothing announces it, so nobody could reach it"
    )
    assert "Saved 'coffee.pdf'" in out


async def test_without_word_the_plain_pdf_renderer_still_produces_a_pdf(design):
    a, out_dir = design
    calls = []

    def _no_word(docx_path, pdf_path):
        raise RuntimeError("Word is not installed")

    def _plain(content, path, title=""):
        calls.append((content, path, title))
        from pathlib import Path
        Path(path).write_bytes(b"%PDF-1.4 plain")

    with patch.object(a, "_docx_to_pdf", _no_word), \
            patch("shared.tools.pdf_render.markdown_to_pdf", _plain):
        out = await a.save_architecture_pdf.ainvoke({"filename": "coffee.pdf"})

    assert calls and "Coffee Ordering App" in calls[0][0]
    assert (out_dir / "coffee.pdf").read_bytes() == b"%PDF-1.4 plain"
    assert sorted(f.name for f in out_dir.iterdir()) == ["coffee.pdf"]
    assert "Saved 'coffee.pdf'" in out


async def test_export_document_as_word_is_the_designed_document(design):
    from docx import Document

    a, out_dir = design
    await a.export_document.ainvoke({"filename": "coffee-export.docx"})

    doc = Document(str(out_dir / "coffee-export.docx"))
    texts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                texts.extend(p.text for p in cell.paragraphs)
    assert "DESIGN DOCUMENT" in "\n".join(texts)


async def test_export_document_as_pdf_uses_the_converted_docx(design):
    a, out_dir = design
    with patch.object(a, "_docx_to_pdf", _fake_convert):
        out = await a.export_document.ainvoke({"filename": "coffee-export.pdf"})

    assert (out_dir / "coffee-export.pdf").read_bytes().startswith(b"%PDF")
    assert "coffee-export.pdf" in out


async def test_other_export_formats_are_untouched(design):
    a, out_dir = design
    await a.export_document.ainvoke({"filename": "notes.md"})
    assert (out_dir / "notes.md").read_text(encoding="utf-8").startswith("# Coffee Ordering App")
