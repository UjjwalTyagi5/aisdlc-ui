"""The migration-intent brief as a designed document — Word (.docx) and PDF.

WHY NOT THE GENERIC MARKDOWN RENDERERS. `shared.tools.doc_export` turns Markdown into a
document with plain headings and grid tables, which is right for prose and wrong for this:
the brief is read by people deciding whether to fund a migration, and at a glance it has to
show what changes, why, what it costs and when. So it is laid out here, from the brief's
structured fields, with the same visual language as the page: a title band running white
into PwC orange, a key-facts strip, the change per part of the system (today in grey with
end of life marked red, the target in orange), the recommendation as a callout, the change
mix as a bar that deepens with how much changes, module cards, a gain/cost table and a
timeline chart.

ONE PLAN, TWO RENDERERS. `_plan()` decides which sections exist and their numbering; the
Word and PDF renderers draw the same sections in the same order, so the two downloads
match each other and the page. The timeline chart is one PNG (Pillow), embedded in both.

Fonts: Word uses Calibri (every Office install has it). The PDF registers the first
TrueType font it finds (DejaVu on Linux, Segoe UI or Arial on Windows) so arrows and dots
render; with none it falls back to Helvetica and plain ASCII stand-ins.
"""
from __future__ import annotations

import io
import os
import textwrap
from datetime import date
from typing import Any, Optional

from agents_orchestrator.requirements_modernization_agent.brief import (
    CHANGE_LABEL,
    DRIVER_LABEL,
    EFFORT_LABEL,
    MILESTONE_LABEL,
    STATUS_LABEL,
    display_drivers,
    display_layers,
    key_facts,
    parse_date,
    pretty_date,
    sorted_milestones,
)
from shared.models.artifacts import MigrationIntentArtifact

# ── the palette (the page's, in print) ──────────────────────────────────────
# One colour, one meaning, the same rules as the page:
#   grey    today, and anything that is context
#   orange  the plan: the target, the recommendation, and how much changes (PwC orange,
#           deepening with the size of the change: upgrade → re-platform → rewrite)
#   red     end of life and risks
#   amber   legacy or support ending (a dot, never a fill)
#   green   fine as it is: a part kept as is, a runtime still supported (never text)
# Orange TEXT is ACCENT (5.2:1 on white, 4.6:1 on TINT_10) or ACCENT_DEEP on a tint
# (6.1:1). The brand orange itself, BRAND, is 3.3:1 on white: a fill, a rule or a mark.

INK = "1F1F1F"
BODY = "3D3D3D"
MUTED = "6B6B6B"
FAINT = "A3A3A3"  # rules, ticks and arrows only: too light for text
RULE = "E5E5E5"
PANEL = "F5F5F5"
WHITE = "FFFFFF"
BRAND = "FD5108"
ACCENT = "C2410C"
ACCENT_DEEP = "9A3412"
TINT_5, TINT_8, TINT_10, TINT_12, TINT_15, TINT_20 = "FFF6F3", "FFF1EB", "FFEEE6", "FFEAE1", "FFE5DA", "FFDCCE"
TINT_30, TINT_45, TINT_60, TINT_80 = "FECBB5", "FEB190", "FE976B", "FD7439"
RED, RED_TINT = "B42318", "FDECEC"
AMBER = "E09A00"
GREEN, GREEN_TINT = "16A34A", "E3F4EA"
GREY_TINT, GREY_INK = "EDEDED", "525252"

STATUS = {  # (dot, label): today is grey, and only end of life is red
    "eol": (RED, RED), "approaching": (AMBER, MUTED), "legacy": (AMBER, MUTED), "supported": (GREEN, MUTED),
}
TARGET = TINT_10
#: Change types from the least to the most that changes: the order of the change mix.
CHANGE_ORDER = ("keep", "upgrade", "replatform", "replace", "rewrite", "new", "retire")
CHANGE = {  # (pill fill, pill text)
    "keep": (GREEN_TINT, INK), "upgrade": (TINT_15, ACCENT_DEEP), "replatform": (TINT_30, "7C2D12"),
    "replace": (TINT_45, "5A1E07"), "rewrite": (ACCENT, WHITE), "new": (ACCENT_DEEP, WHITE),
    "retire": (GREY_TINT, GREY_INK),
}
CHANGE_BAR = {
    "keep": "86D0A4", "upgrade": TINT_30, "replatform": TINT_60, "replace": TINT_80,
    "rewrite": BRAND, "new": ACCENT_DEEP, "retire": "D4D4D4",
}
GAIN = (TINT_8, ACCENT_DEEP)
COST = (PANEL, MUTED)
HEAD_STEPS = (TINT_5, TINT_8, TINT_10, TINT_12, TINT_15)  # a Word header row, white into orange


def _kind_dot(kind: str) -> str:
    """The deadline is the one milestone in orange; the rest are grey."""
    return BRAND if kind == "deadline" else "8C8C8C"


def _kind_ink(kind: str) -> str:
    return ACCENT if kind == "deadline" else MUTED


def _blend(a: str, b: str, f: float) -> str:
    ca, cb = _rgb(a), _rgb(b)
    return "%02X%02X%02X" % tuple(round(ca[i] + (cb[i] - ca[i]) * f) for i in range(3))


def _change_mix(brief: MigrationIntentArtifact) -> list[tuple[str, int]]:
    """(change type, how many modules), from the least to the most that changes."""
    counts: dict[str, int] = {}
    for m in brief.module_changes:
        if m.change_type:
            counts[m.change_type] = counts.get(m.change_type, 0) + 1
    rank = {t: i for i, t in enumerate(CHANGE_ORDER)}
    return sorted(counts.items(), key=lambda kv: rank.get(kv[0], len(rank)))


SECTION_TITLES = {
    "glance": "The change at a glance",
    "why": "Why we are modernizing",
    "recommendation": "Recommended target stack",
    "target_state": "Target state",
    "scope": "Scope",
    "modules": "What changes in each module",
    "tradeoffs": "Trade-offs",
    "timeline": "Timeline",
    "constraints": "Constraints",
    "success": "How we will measure success",
    "people": "Stakeholders",
    "risks": "Assumptions, risks and open questions",
}


def _plan(brief: MigrationIntentArtifact) -> list[str]:
    """The sections this brief has, in order. Empty optional sections are left out."""
    plan = ["glance", "why"]
    rec = brief.recommendation
    if rec and (rec.summary or rec.rationale or rec.alternatives):
        plan.append("recommendation")
    elif brief.target_state.description:
        plan.append("target_state")
    plan.append("scope")
    if brief.module_changes:
        plan.append("modules")
    if brief.trade_offs:
        plan.append("tradeoffs")
    if brief.milestones:
        plan.append("timeline")
    plan += ["constraints", "success"]
    if brief.stakeholders:
        plan.append("people")
    plan.append("risks")
    return plan


def _clean(items) -> list[str]:
    return [str(i).strip() for i in (items or []) if str(i).strip()]


_PLACEHOLDERS = {"", "-", "—", "n/a", "na", "none", "unknown", "tbd", "not measured"}


def _value(text: str) -> str:
    """A measure's value, or "—" for the placeholders a model writes for 'no figure'."""
    return "—" if (text or "").strip().lower() in _PLACEHOLDERS else text.strip()


def _rgb(hex_: str) -> tuple[int, int, int]:
    return tuple(int(hex_[i:i + 2], 16) for i in (0, 2, 4))  # type: ignore[return-value]


# ── fonts ────────────────────────────────────────────────────────────────────

_FONT_CANDIDATES = (
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf"),
    ("C:/Windows/Fonts/segoeui.ttf", "C:/Windows/Fonts/segoeuib.ttf"),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
)


def _font_files() -> Optional[tuple[str, str]]:
    for regular, bold in _FONT_CANDIDATES:
        if os.path.exists(regular) and os.path.exists(bold):
            return regular, bold
    return None


#: For marks the text font lacks (Segoe UI and Arial have no ✓).
_SYMBOL_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "C:/Windows/Fonts/seguisym.ttf", "/Library/Fonts/Apple Symbols.ttf",
)

_PROVIDER = {"ado": "Azure DevOps", "github": "GitHub", "gitlab": "GitLab", "bitbucket": "Bitbucket"}


def repository_line(brief: MigrationIntentArtifact) -> str:
    import urllib.parse  # noqa: PLC0415

    repo = brief.legacy_repository
    if not (repo and (repo.url or repo.name)):
        return "Not named yet."
    provider = _PROVIDER.get((repo.provider or "").lower(), repo.provider)
    where = " / ".join(x for x in (provider, repo.project, repo.name) if x)
    url = urllib.parse.unquote(repo.url or "")
    return (where or url) + (f"  ·  {url}" if url and where else "")


# ── the timeline chart (PNG, used by both documents) ────────────────────────


def timeline_png(brief: MigrationIntentArtifact, width_in: float = 6.85) -> Optional[tuple[bytes, float]]:
    """A horizontal timeline of the dated milestones, or None with fewer than two.
    Returns (png bytes, height in inches at `width_in`)."""
    from PIL import Image, ImageDraw, ImageFont  # noqa: PLC0415

    dated = [(parse_date(m.date), m) for m in sorted_milestones(brief)]
    dated = [(d, m) for d, m in dated if d]
    if len(dated) < 2:
        return None

    scale = 2
    W, H = int(width_in * 150) * scale, 300 * scale
    img = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(img)
    files = _font_files()

    def font(size: int, bold: bool = False):
        if files:
            try:
                return ImageFont.truetype(files[1] if bold else files[0], size * scale)
            except OSError:
                pass
        try:
            return ImageFont.load_default(size * scale)
        except TypeError:  # Pillow < 10.1
            return ImageFont.load_default()

    f_date, f_label, f_tick = font(17, True), font(15), font(13)
    start, end = dated[0][0], dated[-1][0]
    span = max((end - start).days, 1)
    pad = 120 * scale
    axis_y = H // 2

    def x_of(d: date) -> int:
        return int(pad + (d - start).days / span * (W - 2 * pad))

    # The axis runs grey into orange: from where the programme starts to its deadline.
    a0, a1 = pad - 40 * scale, W - pad + 40 * scale
    for x in range(a0, a1, 2 * scale):
        colour = _rgb(_blend(RULE, TINT_45, (x - a0) / (a1 - a0)))
        draw.line([(x, axis_y), (x + 2 * scale, axis_y)], fill=colour, width=6 * scale)
    for year in range(start.year + 1, end.year + 1):  # year ticks
        x = x_of(date(year, 1, 1))
        draw.line([(x, axis_y - 10 * scale), (x, axis_y + 10 * scale)], fill=_rgb(FAINT), width=2 * scale)
        draw.text((x + 4 * scale, axis_y + 12 * scale), str(year), font=f_tick, fill=_rgb(MUTED))
    today = date.today()
    if start <= today <= end:
        x = x_of(today)
        draw.line([(x, axis_y - 26 * scale), (x, axis_y + 26 * scale)], fill=_rgb(MUTED), width=2 * scale)
        draw.text((x - 16 * scale, axis_y + 28 * scale), "Today", font=f_tick, fill=_rgb(MUTED))

    for i, (d, m) in enumerate(dated):
        x = x_of(d)
        colour = _rgb(_kind_dot(m.kind))
        date_colour = _rgb(ACCENT if m.kind == "deadline" else INK)
        up = i % 2 == 0
        stem_end = axis_y - 34 * scale if up else axis_y + 34 * scale
        draw.line([(x, axis_y), (x, stem_end)], fill=colour, width=2 * scale)
        r = 13 * scale
        draw.ellipse([(x - r, axis_y - r), (x + r, axis_y + r)], fill=colour, outline="white", width=4 * scale)
        lines = [pretty_date(m.date)] + textwrap.wrap(m.label, 22)[:2]
        fonts = [f_date] + [f_label] * (len(lines) - 1)
        heights = [draw.textbbox((0, 0), t, font=f)[3] for t, f in zip(lines, fonts)]
        widths = [draw.textlength(t, font=f) for t, f in zip(lines, fonts)]
        block_h = sum(heights) + 4 * scale * (len(lines) - 1)
        top = stem_end - block_h - 6 * scale if up else stem_end + 6 * scale
        for t, f, h, w in zip(lines, fonts, heights, widths):
            left = min(max(x - w / 2, 8 * scale), W - w - 8 * scale)
            draw.text((left, top), t, font=f, fill=date_colour if f is f_date else _rgb(BODY))
            top += h + 4 * scale

    buf = io.BytesIO()
    img.save(buf, "PNG", dpi=(300, 300))
    return buf.getvalue(), width_in * H / W


# ════════════════════════════════════════════════════════════════════════════
#   WORD
# ════════════════════════════════════════════════════════════════════════════


class _Word:
    def __init__(self, brief: MigrationIntentArtifact, meta: dict):
        from docx import Document  # noqa: PLC0415
        from docx.shared import Mm, Pt, RGBColor  # noqa: PLC0415

        self.b, self.meta = brief, meta
        self.Pt, self.Mm, self.RGB = Pt, Mm, RGBColor
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(18)
        sec.top_margin, sec.bottom_margin = Mm(15), Mm(16)
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10)
        normal.font.color.rgb = RGBColor.from_string(BODY)
        self._east_asia(normal.element)
        normal.paragraph_format.space_after = Pt(2)
        normal.paragraph_format.line_spacing = 1.08
        self.doc.core_properties.title = f"Migration-intent brief — {brief.system_name}"
        self.doc.core_properties.subject = "Track 3 · Code Modernization"
        self.n = 0

    # ── XML helpers ──────────────────────────────────────────────────────────
    @staticmethod
    def _el(tag: str, **attrs):
        from docx.oxml import OxmlElement  # noqa: PLC0415
        from docx.oxml.ns import qn  # noqa: PLC0415

        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(f"w:{k}"), str(v))
        return el

    def _east_asia(self, style_el) -> None:
        from docx.oxml.ns import qn  # noqa: PLC0415

        rpr = style_el.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = self._el("w:rFonts")
            rpr.insert(0, fonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attr}"), "Calibri")

    _TCPR_AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
                       "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")

    def shade(self, cell, fill: str) -> None:
        from docx.oxml.ns import qn  # noqa: PLC0415

        tcPr = cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn("w:shd")):
            tcPr.remove(old)
        tcPr.insert_element_before(self._el("w:shd", val="clear", color="auto", fill=fill),
                                   *self._TCPR_AFTER_SHD)

    def pad(self, cell, top=70, bottom=70, left=110, right=110) -> None:
        from docx.oxml.ns import qn  # noqa: PLC0415

        tcPr = cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn("w:tcMar")):
            tcPr.remove(old)
        mar = self._el("w:tcMar")
        for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            mar.append(self._el(f"w:{side}", w=v, type="dxa"))
        tcPr.insert_element_before(mar, *self._TCPR_AFTER_SHD[2:])

    def cell_border(self, cell, **edges) -> None:
        """edges: left=("D2541F", 24) — a coloured bar on one side of a card."""
        from docx.oxml.ns import qn  # noqa: PLC0415

        tcPr = cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn("w:tcBorders")):
            tcPr.remove(old)
        borders = self._el("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            if side in edges:
                colour, size = edges[side]
                borders.append(self._el(f"w:{side}", val="single", sz=size, space=0, color=colour))
        tcPr.insert_element_before(borders, "w:shd", *self._TCPR_AFTER_SHD)

    def table(self, rows: int, widths_mm: list[float], *, rules: str = "none", rule_colour: str = RULE,
              gutter: int = 0):
        """A fixed-layout table. rules: none | rows (horizontal lines) | grid."""
        from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: PLC0415
        from docx.oxml.ns import qn  # noqa: PLC0415

        t = self.doc.add_table(rows=rows, cols=len(widths_mm))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        borders = self._el("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            on = (rules == "grid") or (rules == "rows" and edge in ("insideH", "bottom"))
            if gutter and edge == "insideV":
                borders.append(self._el(f"w:{edge}", val="single", sz=gutter, space=0, color=WHITE))
            elif gutter and edge == "insideH":
                borders.append(self._el(f"w:{edge}", val="single", sz=gutter, space=0, color=WHITE))
            elif on:
                borders.append(self._el(f"w:{edge}", val="single", sz=4, space=0, color=rule_colour))
            else:
                borders.append(self._el(f"w:{edge}", val="nil"))
        tblPr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                                    "w:tblCaption", "w:tblDescription", "w:tblPrChange")
        tblPr.insert_element_before(self._el("w:tblLayout", type="fixed"), "w:tblCellMar", "w:tblLook",
                                    "w:tblCaption", "w:tblDescription", "w:tblPrChange")
        for old in tblPr.findall(qn("w:tblW")):
            old.set(qn("w:type"), "dxa")
            old.set(qn("w:w"), str(int(sum(widths_mm) * 56.7)))
        for row in t.rows:
            for i, w in enumerate(widths_mm):
                row.cells[i].width = self.Mm(w)
                self.pad(row.cells[i])
        return t

    # ── text helpers ─────────────────────────────────────────────────────────
    def run(self, p, text: str, *, size: float = 10, bold=False, italic=False, colour: str = BODY,
            tint: Optional[str] = None, caps=False, spacing: Optional[int] = None):
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        r.font.size = self.Pt(size)
        r.font.color.rgb = self.RGB.from_string(colour)
        if caps:
            r.font.all_caps = True
        rpr = r._r.get_or_add_rPr()
        after_shd = ("w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
                     "w:eastAsianLayout", "w:specVanish", "w:oMath")
        if spacing is not None:  # schema order: spacing comes before sz, shd
            rpr.insert_element_before(self._el("w:spacing", val=spacing), "w:w", "w:kern", "w:position",
                                      "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr",
                                      "w:shd", *after_shd)
        if tint:
            rpr.insert_element_before(self._el("w:shd", val="clear", color="auto", fill=tint), *after_shd)
        return r

    def para(self, container=None, *, before: float = 0, after: float = 2, align=None, keep=False, first=False):
        from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: PLC0415

        container = container if container is not None else self.doc
        if first and getattr(container, "paragraphs", None):
            p = container.paragraphs[0]
        else:
            p = container.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = self.Pt(before), self.Pt(after)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if keep:
            pf.keep_with_next = True
        return p

    def pill(self, p, text: str, tint: str, ink: str, size: float = 8.5) -> None:
        self.run(p, f"\u00a0{text}\u00a0", size=size, bold=True, colour=ink, tint=tint)

    def heading(self, key: str) -> None:
        self.n += 1
        p = self.doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        bdr = self._el("w:pBdr")
        bdr.append(self._el("w:bottom", val="single", sz=6, space=4, color=RULE))
        pPr.append(bdr)
        pf = p.paragraph_format
        pf.space_before, pf.space_after, pf.keep_with_next = self.Pt(16), self.Pt(7), True
        self.run(p, f"{self.n:02d}  ", size=12.5, bold=True, colour=ACCENT)
        self.run(p, SECTION_TITLES[key], size=12.5, bold=True, colour=INK)

    def subhead(self, text: str, container=None, colour: str = INK, before: float = 6) -> None:
        p = self.para(container, before=before, after=2, keep=True)
        self.run(p, text, size=9.5, bold=True, colour=colour)

    def bullets(self, container, items: list[str], mark: str = "•", mark_colour: str = ACCENT,
                size: float = 9.5, first: bool = False, empty: str = "None recorded.") -> None:
        items = _clean(items) or [empty]
        for i, item in enumerate(items):
            p = self.para(container, after=1.5, first=first and i == 0)
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = self.Mm(4), self.Mm(-4)
            self.run(p, f"{mark}  ", size=size, bold=True, colour=mark_colour)
            self.run(p, item, size=size, colour=BODY)

    # ── sections ─────────────────────────────────────────────────────────────
    def title_band(self) -> None:
        """A strip running into PwC orange over a warm band. Word has no gradient fill
        for a table cell, so the strip is 24 cells stepping through the ramp."""
        from docx.enum.table import WD_ROW_HEIGHT_RULE  # noqa: PLC0415

        b = self.b
        steps = 24
        t = self.table(2, [174 / steps] * steps)
        strip = t.rows[0]
        strip.height, strip.height_rule = self.Mm(1.8), WD_ROW_HEIGHT_RULE.EXACTLY
        for i, cell in enumerate(strip.cells):
            self.shade(cell, _blend(TINT_15, BRAND, i / (steps - 1)))
            self.pad(cell, top=0, bottom=0, left=0, right=0)
        cell = t.rows[1].cells[0].merge(t.rows[1].cells[-1])
        self.shade(cell, TINT_8)
        self.pad(cell, top=300, bottom=300, left=340, right=340)
        p = self.para(cell, after=4, first=True)
        self.run(p, "MIGRATION-INTENT BRIEF", size=8, bold=True, colour=ACCENT_DEEP, spacing=30)
        self.run(p, "   ·   TRACK 3 · CODE MODERNIZATION", size=8, colour=MUTED, spacing=20)
        p = self.para(cell, after=4)
        self.run(p, b.system_name or "Unnamed system", size=24, bold=True, colour=INK)
        if b.goal:
            p = self.para(cell, after=6)
            self.run(p, b.goal.strip(), size=11, colour=BODY)
        p = self.para(cell, after=0)
        bits = []
        if b.recorded_at:
            bits.append(f"Recorded {pretty_date(b.recorded_at[:10])}")
        if self.meta.get("version"):
            bits.append(f"Version {self.meta['version']}")
        if self.meta.get("status"):
            bits.append(str(self.meta["status"]).capitalize())
        self.run(p, "   ·   ".join(bits) or "Draft", size=8.5, colour=MUTED)

    def facts(self) -> None:
        facts = key_facts(self.b)
        if not facts:
            return
        self.para(after=2)
        width = 174 / len(facts)
        t = self.table(1, [width] * len(facts), gutter=36)
        for i, (label, value) in enumerate(facts):
            cell = t.rows[0].cells[i]
            self.shade(cell, PANEL)
            self.pad(cell, top=110, bottom=110, left=150, right=110)
            p = self.para(cell, after=1, first=True)
            self.run(p, label, size=7.5, bold=True, colour=MUTED, caps=True, spacing=10)
            p = self.para(cell, after=0)
            colour = RED if label.startswith("End of life") else INK
            self.run(p, value, size=11.5 if len(value) < 24 else 9.5, bold=True, colour=colour)

    def _head_row(self, row, labels: list[str]) -> None:
        """A header row running white into orange, one tint step per column."""
        last = len(HEAD_STEPS) - 1
        for i, head in enumerate(labels):
            cell = row.cells[i]
            self.shade(cell, HEAD_STEPS[round(i * last / max(len(labels) - 1, 1))])
            p = self.para(cell, first=True, after=0)
            self.run(p, head, size=7.5, bold=True, colour=ACCENT_DEEP, spacing=10)

    def _today(self, cell, current: str, status: str, size: float = 9.5) -> None:
        """Today in grey; end of life adds a red edge and a red label."""
        self.shade(cell, PANEL)
        if status == "eol":
            self.cell_border(cell, left=(RED, 18))
        p = self.para(cell, first=True, after=0)
        self.run(p, current or "—", size=size, bold=True, colour=INK)
        if status in STATUS:
            dot, ink = STATUS[status]
            p = self.para(cell, after=0)
            self.run(p, "● ", size=7.5, bold=True, colour=dot)
            self.run(p, STATUS_LABEL[status], size=7.5, bold=True, colour=ink)

    def _target(self, cell, target: str, size: float = 9.5) -> None:
        self.shade(cell, TARGET)
        p = self.para(cell, first=True, after=0)
        self.run(p, target or "—", size=size, bold=True, colour=INK)

    def effort(self, p, effort: str) -> None:
        level = {"low": 1, "medium": 2, "high": 3}.get(effort, 0)
        self.run(p, "●" * level, size=8, colour=BRAND)
        self.run(p, "●" * (3 - level) + " ", size=8, colour="DDDDDD")
        self.run(p, EFFORT_LABEL.get(effort, effort), size=8, bold=True, colour=MUTED)

    def glance(self) -> None:
        self.heading("glance")
        layers = display_layers(self.b)
        if not layers:
            self.bullets(self.doc, [], empty="Not recorded yet.")
            return
        t = self.table(len(layers) + 1, [40, 50, 8, 50, 26], rules="rows")
        self._head_row(t.rows[0], ["PART OF THE SYSTEM", "TODAY", "", "TARGET", "CHANGE"])
        for r, layer in enumerate(layers, start=1):
            cells = t.rows[r].cells
            p = self.para(cells[0], first=True, after=0)
            self.run(p, layer.layer, size=10, bold=True, colour=INK)
            if layer.modules:
                p = self.para(cells[0], after=0)
                self.run(p, ", ".join(layer.modules), size=7.5, colour=MUTED)
            self._today(cells[1], layer.current, layer.current_status)
            p = self.para(cells[2], first=True, after=0, align="center")
            self.run(p, "→", size=14, bold=True, colour=FAINT)
            self._target(cells[3], layer.target)
            p = self.para(cells[4], first=True, after=0)
            if layer.change_type:
                self.pill(p, CHANGE_LABEL[layer.change_type], *CHANGE[layer.change_type])
            for c in cells:
                self._valign(c)

    def _valign(self, cell) -> None:
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT  # noqa: PLC0415

        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    def why(self) -> None:
        self.heading("why")
        drivers = display_drivers(self.b)
        if not drivers:
            self.bullets(self.doc, [])
            return
        rows = (len(drivers) + 1) // 2
        t = self.table(rows, [87, 87], gutter=60)
        for i, d in enumerate(drivers):
            cell = t.rows[i // 2].cells[i % 2]
            self.shade(cell, PANEL)
            self.cell_border(cell, left=(BRAND, 24))
            self.pad(cell, top=110, bottom=110, left=170, right=130)
            p = self.para(cell, first=True, after=1)
            self.run(p, DRIVER_LABEL.get(d.category, "Other"), size=7.5, bold=True, colour=MUTED, caps=True, spacing=10)
            p = self.para(cell, after=1)
            self.run(p, d.title or d.detail, size=10, bold=True, colour=INK)
            if d.detail and d.title:
                p = self.para(cell, after=0)
                self.run(p, d.detail, size=9, colour=BODY)
        if len(drivers) % 2:
            self.shade(t.rows[-1].cells[1], WHITE)

    def recommendation(self) -> None:
        self.heading("recommendation")
        rec = self.b.recommendation
        t = self.table(1, [174])
        cell = t.rows[0].cells[0]
        self.shade(cell, TINT_8)
        self.cell_border(cell, left=(BRAND, 36))
        self.pad(cell, top=140, bottom=140, left=220, right=180)
        p = self.para(cell, first=True, after=3)
        by = ("RECOMMENDED BY THE REQUIREMENTS AGENT" if rec.recommended_by == "agent"
              else "SET BY THE BUSINESS")
        self.run(p, by, size=7.5, bold=True, colour=ACCENT, spacing=20)
        if rec.summary:
            p = self.para(cell, after=2)
            self.run(p, rec.summary.strip(), size=10.5, colour=INK)
        if rec.recommended_by == "agent":
            p = self.para(cell, after=0)
            self.run(p, "Accepted when this brief is signed off.", size=8, italic=True, colour=MUTED)
        if rec.rationale:
            self.subhead("Why this stack", before=8)
            self.bullets(self.doc, rec.rationale, mark="✓", mark_colour=ACCENT)
        if rec.alternatives:
            self.subhead("Alternatives considered", before=8)
            t = self.table(len(rec.alternatives) + 1, [58, 116], rules="rows")
            for i, head in enumerate(["OPTION", "WHY NOT"]):
                self.shade(t.rows[0].cells[i], PANEL)
                p = self.para(t.rows[0].cells[i], first=True, after=0)
                self.run(p, head, size=7.5, bold=True, colour=MUTED, spacing=10)
            for r, alt in enumerate(rec.alternatives, start=1):
                p = self.para(t.rows[r].cells[0], first=True, after=0)
                self.run(p, alt.option, size=9.5, bold=True, colour=INK)
                p = self.para(t.rows[r].cells[1], first=True, after=0)
                self.run(p, alt.why_not or "—", size=9.5, colour=BODY)

    def target_state(self) -> None:
        self.heading("target_state")
        p = self.para(after=2)
        self.run(p, self.b.target_state.description.strip(), size=10, colour=BODY)

    def scope(self) -> None:
        self.heading("scope")
        t = self.table(2, [87, 87], gutter=60)
        for i, (label, tint, ink) in enumerate((("IN SCOPE", TINT_15, ACCENT_DEEP),
                                               ("OUT OF SCOPE", GREY_TINT, GREY_INK))):
            head = t.rows[0].cells[i]
            self.shade(head, tint)
            p = self.para(head, first=True, after=0)
            self.run(p, label, size=7.5, bold=True, colour=ink, spacing=20)
        self.bullets(t.rows[1].cells[0], self.b.in_scope, mark="✓", mark_colour=ACCENT, first=True)
        self.bullets(t.rows[1].cells[1], self.b.out_of_scope, mark="–", mark_colour=MUTED, first=True,
                     empty="Nothing excluded.")
        for c in t.rows[1].cells:
            self.shade(c, PANEL)

    def change_mix(self) -> None:
        """How much of the system changes: a bar from the lightest orange (upgrade) to the
        strongest (rewrite), with its legend."""
        from docx.enum.table import WD_ROW_HEIGHT_RULE  # noqa: PLC0415

        mix = _change_mix(self.b)
        total = sum(n for _, n in mix)
        if not total:
            return
        t = self.table(1, [174 * n / total for _, n in mix], gutter=24)
        row = t.rows[0]
        row.height, row.height_rule = self.Mm(2.6), WD_ROW_HEIGHT_RULE.EXACTLY
        for cell, (kind, _) in zip(row.cells, mix):
            self.shade(cell, CHANGE_BAR.get(kind, "D4D4D4"))
            self.pad(cell, top=0, bottom=0, left=0, right=0)
        p = self.para(before=3, after=6)
        for kind, n in mix:
            self.run(p, "■ ", size=9, colour=CHANGE_BAR.get(kind, "D4D4D4"))
            self.run(p, f"{n} {CHANGE_LABEL.get(kind, kind).lower()}      ", size=8.5, colour=MUTED)

    def modules(self) -> None:
        self.heading("modules")
        mods = self.b.module_changes
        self.change_mix()
        t = self.table(len(mods) + 1, [46, 36, 44, 24, 24], rules="rows")
        self._head_row(t.rows[0], ["MODULE", "TODAY", "TARGET", "CHANGE", "EFFORT"])
        for r, m in enumerate(mods, start=1):
            cells = t.rows[r].cells
            p = self.para(cells[0], first=True, after=0)
            self.run(p, m.module, size=9.5, bold=True, colour=INK)
            if m.path and m.path != m.module:
                p = self.para(cells[0], after=0)
                self.run(p, m.path, size=7.5, colour=MUTED)
            self._today(cells[1], m.current, m.current_status, size=9)
            self._target(cells[2], m.target, size=9)
            p = self.para(cells[3], first=True, after=0)
            if m.change_type:
                self.pill(p, CHANGE_LABEL[m.change_type], *CHANGE[m.change_type])
            p = self.para(cells[4], first=True, after=0)
            if m.effort:
                self.effort(p, m.effort)
            for c in cells:
                self._valign(c)
        detailed = [m for m in mods if _clean(m.changes)]
        if detailed:
            self.subhead("What changes", before=10)
            for m in detailed:
                card = self.table(1, [174])
                cell = card.rows[0].cells[0]
                self.shade(cell, PANEL)
                self.cell_border(cell, left=(CHANGE_BAR.get(m.change_type, "D4D4D4"), 18))
                self.pad(cell, top=90, bottom=90, left=170, right=130)
                p = self.para(cell, first=True, after=2)
                self.run(p, m.module + "  ", size=10, bold=True, colour=INK)
                if m.change_type:
                    self.pill(p, CHANGE_LABEL[m.change_type], *CHANGE[m.change_type], size=7.5)
                if m.current or m.target:
                    self.run(p, f"   {m.current or '—'}  →  {m.target or '—'}", size=8.5, colour=MUTED)
                self.bullets(cell, m.changes, mark="•", mark_colour=ACCENT, size=9)
                self.para(after=1)

    def tradeoffs(self) -> None:
        self.heading("tradeoffs")
        items = self.b.trade_offs
        t = self.table(len(items) + 1, [50, 62, 62], rules="rows")
        for i, (head, colour) in enumerate((("DECISION", MUTED), ("WHAT WE GAIN", GAIN[1]),
                                            ("WHAT IT COSTS", COST[1]))):
            p = self.para(t.rows[0].cells[i], first=True, after=0)
            self.run(p, head, size=7.5, bold=True, colour=colour, spacing=10)
        for r, item in enumerate(items, start=1):
            cells = t.rows[r].cells
            p = self.para(cells[0], first=True, after=0)
            self.run(p, item.decision, size=9.5, bold=True, colour=INK)
            self.shade(cells[1], GAIN[0])
            p = self.para(cells[1], first=True, after=0)
            self.run(p, item.gain or "—", size=9, colour=BODY)
            self.shade(cells[2], COST[0])
            p = self.para(cells[2], first=True, after=0)
            self.run(p, item.cost or "—", size=9, colour=BODY)

    def timeline(self) -> None:
        from docx.shared import Inches  # noqa: PLC0415

        self.heading("timeline")
        chart = timeline_png(self.b)
        if chart:
            png, _h = chart
            p = self.para(after=4, align="center")
            p.add_run().add_picture(io.BytesIO(png), width=Inches(6.85))
        items = sorted_milestones(self.b)
        t = self.table(len(items), [34, 110, 30], rules="rows")
        for r, m in enumerate(items):
            cells = t.rows[r].cells
            p = self.para(cells[0], first=True, after=0)
            self.run(p, "● ", size=9, bold=True, colour=_kind_dot(m.kind))
            self.run(p, pretty_date(m.date), size=9.5, bold=True, colour=INK)
            p = self.para(cells[1], first=True, after=0)
            self.run(p, m.label, size=9.5, colour=BODY)
            p = self.para(cells[2], first=True, after=0)
            self.run(p, MILESTONE_LABEL.get(m.kind, "Milestone"), size=8, bold=True, colour=_kind_ink(m.kind))

    def constraints(self) -> None:
        self.heading("constraints")
        self.bullets(self.doc, self.b.constraints, mark="■", mark_colour=ACCENT, size=9.5)

    def success(self) -> None:
        self.heading("success")
        measures = self.b.success_measures
        if measures:
            t = self.table(len(measures) + 1, [80, 44, 50], rules="rows")
            for i, head in enumerate(["MEASURE", "TODAY", "TARGET"]):
                self.shade(t.rows[0].cells[i], PANEL)
                p = self.para(t.rows[0].cells[i], first=True, after=0)
                self.run(p, head, size=7.5, bold=True, colour=MUTED, spacing=10)
            for r, m in enumerate(measures, start=1):
                cells = t.rows[r].cells
                p = self.para(cells[0], first=True, after=0)
                self.run(p, m.metric, size=9.5, bold=True, colour=INK)
                p = self.para(cells[1], first=True, after=0)
                self.run(p, _value(m.current), size=9.5, colour=MUTED)
                self.shade(cells[2], TARGET)
                p = self.para(cells[2], first=True, after=0)
                self.run(p, m.target or "—", size=10, bold=True, colour=ACCENT)
            self.subhead("Acceptance criteria", before=8)
        self.bullets(self.doc, self.b.success_criteria, mark="○", mark_colour=ACCENT)

    def people(self) -> None:
        self.heading("people")
        people = self.b.stakeholders
        rows = (len(people) + 1) // 2
        t = self.table(rows, [87, 87], gutter=60)
        for i, s in enumerate(people):
            cell = t.rows[i // 2].cells[i % 2]
            self.shade(cell, PANEL)
            p = self.para(cell, first=True, after=0)
            self.run(p, s.name, size=10, bold=True, colour=INK)
            if s.role:
                p = self.para(cell, after=0)
                self.run(p, s.role, size=8.5, colour=MUTED)
        if len(people) % 2:
            self.shade(t.rows[-1].cells[1], WHITE)

    def risks(self) -> None:
        self.heading("risks")
        t = self.table(2, [58, 58, 58], gutter=48)
        cols = (("ASSUMPTIONS", self.b.assumptions, GREY_TINT, GREY_INK, MUTED, "None recorded."),
                ("RISKS", self.b.risks, RED_TINT, RED, RED, "None recorded."),
                ("OPEN QUESTIONS", self.b.open_questions, GREY_TINT, GREY_INK, MUTED, "None."))
        for i, (label, items, tint, ink, mark, empty) in enumerate(cols):
            head = t.rows[0].cells[i]
            self.shade(head, tint)
            p = self.para(head, first=True, after=0)
            self.run(p, label, size=7.5, bold=True, colour=ink, spacing=20)
            body = t.rows[1].cells[i]
            self.shade(body, PANEL)
            self.bullets(body, items, mark="•", mark_colour=mark, size=9, first=True, empty=empty)

    def repository(self) -> None:
        p = self.para(before=14, after=0)
        self.run(p, "LEGACY REPOSITORY   ", size=7.5, bold=True, colour=MUTED, spacing=10)
        self.run(p, repository_line(self.b), size=8.5, colour=BODY)

    def footer(self) -> None:
        sec = self.doc.sections[0]
        p = sec.footer.paragraphs[0]
        self.run(p, f"{self.b.system_name or 'Migration'} · Migration-intent brief · Page ", size=8, colour=MUTED)
        r = p.add_run()
        r.font.size = self.Pt(8)
        r.font.color.rgb = self.RGB.from_string(MUTED)
        from docx.oxml.ns import qn  # noqa: PLC0415

        begin = self._el("w:fldChar", fldCharType="begin")
        instr = self._el("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        end = self._el("w:fldChar", fldCharType="end")
        r._r.append(begin)
        r._r.append(instr)
        r._r.append(end)

    def render(self, path: str) -> None:
        self.title_band()
        self.facts()
        for key in _plan(self.b):
            getattr(self, key)()
        self.repository()
        self.footer()
        self.doc.save(path)


# ════════════════════════════════════════════════════════════════════════════
#   PDF
# ════════════════════════════════════════════════════════════════════════════


class _Pdf:
    def __init__(self, brief: MigrationIntentArtifact, meta: dict):
        from reportlab.lib.pagesizes import A4  # noqa: PLC0415
        from reportlab.lib.units import mm  # noqa: PLC0415

        self.b, self.meta = brief, meta
        self.mm = mm
        self.page = A4
        self.usable = A4[0] - 36 * mm
        self.font, self.bold = self._register_fonts()
        self.n = 0
        self.story: list[Any] = []

    def _register_fonts(self) -> tuple[str, str]:
        from reportlab.pdfbase import pdfmetrics  # noqa: PLC0415
        from reportlab.pdfbase.ttfonts import TTFont  # noqa: PLC0415

        files = _font_files()
        if files:
            try:
                if "BriefSans" not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont("BriefSans", files[0]))
                    pdfmetrics.registerFont(TTFont("BriefSans-Bold", files[1]))
                    from reportlab.lib.fonts import addMapping  # noqa: PLC0415

                    addMapping("BriefSans", 0, 0, "BriefSans")
                    addMapping("BriefSans", 1, 0, "BriefSans-Bold")
                    addMapping("BriefSans", 0, 1, "BriefSans")
                    addMapping("BriefSans", 1, 1, "BriefSans-Bold")
                self.unicode = True
                self.cmap = set(pdfmetrics.getFont("BriefSans").face.charToGlyph)
                self.symbol = None
                for path in _SYMBOL_CANDIDATES:
                    if os.path.exists(path):
                        if "BriefSymbol" not in pdfmetrics.getRegisteredFontNames():
                            pdfmetrics.registerFont(TTFont("BriefSymbol", path))
                        self.symbol = "BriefSymbol"
                        break
                return "BriefSans", "BriefSans-Bold"
            except Exception:  # noqa: BLE001 — a broken font file is not a reason to fail the export
                pass
        self.unicode, self.cmap, self.symbol = False, set(), None
        return "Helvetica", "Helvetica-Bold"

    _STAND_IN = {"→": "»", "●": "•", "✓": "+", "○": "-", "■": "•", "–": "-"}

    def g(self, ch: str) -> str:
        """A glyph, or its ASCII stand-in when no registered font has it."""
        if self.unicode and (ord(ch) in self.cmap or self.symbol):
            return ch
        return self._STAND_IN.get(ch, ch)

    def mark_font(self, ch: str) -> str:
        """The font a bullet mark is drawn in: the text font when it has the glyph."""
        if self.unicode and ord(ch) not in self.cmap and self.symbol:
            return self.symbol
        return self.bold

    # ── primitives ───────────────────────────────────────────────────────────
    def style(self, size=9.5, colour=BODY, bold=False, leading=None, align=0, after=0, before=0):
        from reportlab.lib import colors  # noqa: PLC0415
        from reportlab.lib.styles import ParagraphStyle  # noqa: PLC0415

        return ParagraphStyle(
            "s", fontName=self.bold if bold else self.font, fontSize=size,
            leading=leading or size * 1.32, textColor=colors.HexColor("#" + colour),
            alignment=align, spaceAfter=after, spaceBefore=before,
        )

    def P(self, text: str, **kw):
        from xml.sax.saxutils import escape  # noqa: PLC0415

        from reportlab.platypus import Paragraph  # noqa: PLC0415

        return Paragraph(escape(text or ""), self.style(**kw))

    def H(self, markup: str, **kw):
        from reportlab.platypus import Paragraph  # noqa: PLC0415

        return Paragraph(markup, self.style(**kw))

    @staticmethod
    def esc(text: str) -> str:
        from xml.sax.saxutils import escape  # noqa: PLC0415

        return escape(text or "")

    def pill(self, text: str, tint: str, ink: str, size: float = 7.5) -> str:
        return (f'<font name="{self.bold}" size="{size}" color="#{ink}" backColor="#{tint}">'
                f"&nbsp;{self.esc(text)}&nbsp;</font>")

    def C(self, hex_: str):
        from reportlab.lib import colors  # noqa: PLC0415

        return colors.HexColor("#" + hex_)

    def table(self, data, widths_mm, style_cmds, *, repeat=0):
        from reportlab.platypus import Table, TableStyle  # noqa: PLC0415

        t = Table(data, colWidths=[w * self.mm for w in widths_mm], repeatRows=repeat)
        base = [
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]
        t.setStyle(TableStyle(base + style_cmds))
        return t

    def heading(self, key: str):
        from reportlab.platypus import Spacer  # noqa: PLC0415
        from reportlab.platypus.flowables import HRFlowable  # noqa: PLC0415

        self.n += 1
        return [
            Spacer(1, 12),
            self.H(f'<font color="#{ACCENT}">{self.n:02d}</font>&nbsp;&nbsp;{self.esc(SECTION_TITLES[key])}',
                   size=12.5, colour=INK, bold=True),
            HRFlowable(width="100%", thickness=0.6, color=self.C(RULE), spaceBefore=3, spaceAfter=6),
        ]

    def bullets(self, items, mark="•", mark_colour=ACCENT, size=9.2, empty="None recorded."):
        from reportlab.platypus import Paragraph  # noqa: PLC0415

        out = []
        for item in _clean(items) or [empty]:
            st = self.style(size=size, after=2)
            st.leftIndent, st.bulletIndent = 11, 0
            st.bulletFontName, st.bulletFontSize = self.mark_font(mark), size
            st.bulletColor = self.C(mark_colour)
            out.append(Paragraph(self.esc(item), st, bulletText=self.g(mark)))
        return out

    def section(self, key: str, body: list):
        """A heading never ends a page on its own: it starts a new page unless there is
        room for it and the first ~120pt of its content. Long tables then continue onto
        the next page (header row repeated) instead of leaving half a page blank."""
        from reportlab.platypus import CondPageBreak  # noqa: PLC0415

        self.story.append(CondPageBreak(150))
        self.story += self.heading(key) + body

    # ── sections ─────────────────────────────────────────────────────────────
    def glyph(self, ch: str, colour: str) -> str:
        """A mark in its colour, in a font that has it."""
        return f'<font name="{self.mark_font(ch)}" color="#{colour}">{self.g(ch)}</font>'

    def _head(self, labels) -> tuple[list, list]:
        """A header row running white into orange."""
        row = [self.P(h, size=7, colour=ACCENT_DEEP, bold=True) for h in labels]
        return row, [("BACKGROUND", (0, 0), (-1, 0), ["HORIZONTAL", self.C(TINT_5), self.C(TINT_15)])]

    def title_band(self):
        b = self.b
        rows = [[self.H(f'<font color="#{ACCENT_DEEP}" name="{self.bold}">MIGRATION-INTENT BRIEF</font>'
                        f'<font color="#{MUTED}">&nbsp;&nbsp;·&nbsp;&nbsp;TRACK 3 · CODE MODERNIZATION</font>', size=7.5)],
                [self.P(b.system_name or "Unnamed system", size=23, colour=INK, bold=True, leading=28)]]
        if b.goal:
            rows.append([self.P(b.goal.strip(), size=10.5, colour=BODY, leading=14.5)])
        bits = []
        if b.recorded_at:
            bits.append(f"Recorded {pretty_date(b.recorded_at[:10])}")
        if self.meta.get("version"):
            bits.append(f"Version {self.meta['version']}")
        if self.meta.get("status"):
            bits.append(str(self.meta["status"]).capitalize())
        rows.append([self.P("   ·   ".join(bits) or "Draft", size=8, colour=MUTED)])
        t = self.table(rows, [174], [
            # White into PwC orange, as on the page.
            ("BACKGROUND", (0, 0), (-1, -1), ["HORIZONTAL", self.C(WHITE), self.C(TINT_5), self.C(TINT_20)]),
            ("BOX", (0, 0), (-1, -1), 0.6, self.C(TINT_20)),
            ("ROUNDEDCORNERS", [7, 7, 7, 7]),
            ("LEFTPADDING", (0, 0), (-1, -1), 16), ("RIGHTPADDING", (0, 0), (-1, -1), 16),
            ("TOPPADDING", (0, 0), (-1, 0), 14), ("BOTTOMPADDING", (0, -1), (-1, -1), 14),
            ("TOPPADDING", (0, 1), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -2), 3),
        ])
        self.story.append(t)

    def facts(self):
        from reportlab.platypus import Spacer  # noqa: PLC0415

        facts = key_facts(self.b)
        if not facts:
            return
        w = 174 / len(facts)
        cells = []
        for label, value in facts:
            colour = RED if label.startswith("End of life") else INK
            cells.append([self.P(label.upper(), size=6.8, colour=MUTED, bold=True),
                          self.P(value, size=11 if len(value) < 24 else 8.8, colour=colour, bold=True, leading=13)])
        t = self.table([cells], [w] * len(facts), [
            ("BACKGROUND", (0, 0), (-1, -1), self.C(PANEL)),
            ("LINEAFTER", (0, 0), (-2, -1), 3, self.C(WHITE)),
            ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ("LEFTPADDING", (0, 0), (-1, -1), 9),
        ])
        self.story += [Spacer(1, 6), t]

    def _today_cell(self, current: str, status: str):
        parts = [self.P(current or "—", size=8.8, colour=INK, bold=True, leading=11)]
        if status in STATUS:
            dot, ink = STATUS[status]
            parts.append(self.H(f'{self.glyph("●", dot)} <font color="#{ink}" name="{self.bold}">'
                                f"{STATUS_LABEL[status]}</font>", size=7))
        return parts

    def _today_cmds(self, col: int, row: int, status: str) -> list:
        """Today in grey; end of life adds a red edge."""
        cmds = [("BACKGROUND", (col, row), (col, row), self.C(PANEL))]
        if status == "eol":
            cmds.append(("LINEBEFORE", (col, row), (col, row), 2.5, self.C(RED)))
        return cmds

    def glance(self):
        layers = display_layers(self.b)
        if not layers:
            self.section("glance", self.bullets([], empty="Not recorded yet."))
            return
        head, cmds = self._head(("PART OF THE SYSTEM", "TODAY", "", "TARGET", "CHANGE"))
        data = [head]
        cmds.append(("LINEBELOW", (0, 1), (-1, -1), 0.5, self.C(RULE)))
        for r, layer in enumerate(layers, start=1):
            first = [self.P(layer.layer, size=9.5, colour=INK, bold=True, leading=12)]
            if layer.modules:
                first.append(self.P(", ".join(layer.modules), size=6.8, colour=MUTED))
            change = (self.H(self.pill(CHANGE_LABEL[layer.change_type], *CHANGE[layer.change_type]), size=8)
                      if layer.change_type else "")
            data.append([first, self._today_cell(layer.current, layer.current_status),
                         self.H(f'<font color="#{FAINT}" name="{self.bold}">{self.g("→")}</font>', size=13, align=1),
                         self.P(layer.target or "—", size=8.8, colour=INK, bold=True, leading=11), change])
            cmds += self._today_cmds(1, r, layer.current_status) + [("BACKGROUND", (3, r), (3, r), self.C(TARGET))]
        self.section("glance", [self.table(data, [40, 50, 8, 50, 26], cmds, repeat=1)])

    def why(self):
        drivers = display_drivers(self.b)
        if not drivers:
            self.section("why", self.bullets([]))
            return
        cells = []
        for d in drivers:
            card = [self.P(DRIVER_LABEL.get(d.category, "Other").upper(), size=6.8, colour=MUTED, bold=True),
                    self.P(d.title or d.detail, size=9.5, colour=INK, bold=True, leading=12)]
            if d.detail and d.title:
                card.append(self.P(d.detail, size=8.5, colour=BODY, leading=11))
            cells.append(card)
        rows, cmds = [], []
        for i in range(0, len(cells), 2):
            pair = cells[i:i + 2]
            rows.append([pair[0], "", pair[1] if len(pair) > 1 else ""])
            r = i // 2
            for j in range(len(pair)):
                cmds += [("BACKGROUND", (j * 2, r), (j * 2, r), self.C(PANEL)),
                         ("LINEBEFORE", (j * 2, r), (j * 2, r), 2.5, self.C(BRAND))]
        cmds += [("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 7),
                 ("BOTTOMPADDING", (0, 0), (-1, -1), 7), ("LEFTPADDING", (0, 0), (-1, -1), 9)]
        grid = self.table(rows, [85, 4, 85], cmds)
        grid.setStyle([("LINEBELOW", (0, 0), (-1, -2), 5, self.C(WHITE))])
        self.section("why", [grid])

    def recommendation(self):
        from reportlab.platypus import Spacer  # noqa: PLC0415

        rec = self.b.recommendation
        by = "RECOMMENDED BY THE REQUIREMENTS AGENT" if rec.recommended_by == "agent" else "SET BY THE BUSINESS"
        inner = [self.P(by, size=6.8, colour=ACCENT, bold=True)]
        if rec.summary:
            inner.append(self.P(rec.summary.strip(), size=10, colour=INK, leading=14))
        if rec.recommended_by == "agent":
            inner.append(self.P("Accepted when this brief is signed off.", size=7.5, colour=MUTED))
        callout = self.table([[inner]], [174], [
            ("BACKGROUND", (0, 0), (-1, -1), self.C(TINT_8)),
            ("LINEBEFORE", (0, 0), (0, -1), 3.5, self.C(BRAND)),
            ("LEFTPADDING", (0, 0), (-1, -1), 12), ("TOPPADDING", (0, 0), (-1, -1), 9),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ])
        body: list = [callout]
        if rec.rationale:
            body += [Spacer(1, 6), self.P("Why this stack", size=9.5, colour=INK, bold=True, after=3)]
            body += self.bullets(rec.rationale, mark="✓", mark_colour=ACCENT)
        if rec.alternatives:
            body += [Spacer(1, 6), self.P("Alternatives considered", size=9.5, colour=INK, bold=True, after=3)]
            data = [[self.P("OPTION", size=6.8, colour=MUTED, bold=True), self.P("WHY NOT", size=6.8, colour=MUTED, bold=True)]]
            data += [[self.P(a.option, size=8.8, colour=INK, bold=True), self.P(a.why_not or "—", size=8.8)]
                     for a in rec.alternatives]
            body.append(self.table(data, [58, 116], [("BACKGROUND", (0, 0), (-1, 0), self.C(PANEL)),
                                                     ("LINEBELOW", (0, 1), (-1, -1), 0.5, self.C(RULE))]))
        self.section("recommendation", body)

    def target_state(self):
        self.section("target_state", [self.P(self.b.target_state.description.strip(), size=9.5)])

    def scope(self):
        data = [[self.P("IN SCOPE", size=6.8, colour=ACCENT_DEEP, bold=True), "",
                 self.P("OUT OF SCOPE", size=6.8, colour=GREY_INK, bold=True)],
                [self.bullets(self.b.in_scope, mark="✓", mark_colour=ACCENT), "",
                 self.bullets(self.b.out_of_scope, mark="–", mark_colour=MUTED, empty="Nothing excluded.")]]
        t = self.table(data, [85, 4, 85], [
            ("BACKGROUND", (0, 0), (0, 0), self.C(TINT_15)), ("BACKGROUND", (2, 0), (2, 0), self.C(GREY_TINT)),
            ("BACKGROUND", (0, 1), (0, 1), self.C(PANEL)), ("BACKGROUND", (2, 1), (2, 1), self.C(PANEL)),
            ("VALIGN", (0, 1), (-1, 1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 9),
        ])
        self.section("scope", [t])

    def change_mix(self) -> list:
        """How much of the system changes: a bar from the lightest orange (upgrade) to the
        strongest (rewrite), with its legend."""
        from reportlab.platypus import Spacer, Table, TableStyle  # noqa: PLC0415

        mix = _change_mix(self.b)
        total = sum(n for _, n in mix)
        if not total:
            return []
        bar = Table([[""] * len(mix)], colWidths=[174 * n / total * self.mm for _, n in mix],
                    rowHeights=[3 * self.mm])
        cmds = [("BACKGROUND", (i, 0), (i, 0), self.C(CHANGE_BAR.get(k, "D4D4D4"))) for i, (k, _) in enumerate(mix)]
        cmds += [("ROUNDEDCORNERS", [4, 4, 4, 4]), ("TOPPADDING", (0, 0), (-1, -1), 0),
                 ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]
        if len(mix) > 1:  # a 2pt gap keeps neighbouring steps of the ramp apart
            cmds.append(("LINEAFTER", (0, 0), (-2, 0), 2, self.C(WHITE)))
        bar.setStyle(TableStyle(cmds))
        legend = "&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;".join(
            f'{self.glyph("■", CHANGE_BAR.get(k, "D4D4D4"))}&nbsp;{n} {self.esc(CHANGE_LABEL.get(k, k).lower())}'
            for k, n in mix)
        return [bar, Spacer(1, 4), self.H(legend, size=8, colour=MUTED), Spacer(1, 7)]

    def effort(self, effort: str) -> str:
        level = {"low": 1, "medium": 2, "high": 3}.get(effort, 0)
        dots = self.glyph("●", BRAND) * level + self.glyph("●", "DDDDDD") * (3 - level)
        return f'{dots}&nbsp;<font name="{self.bold}" color="#{MUTED}" size="7.5">{EFFORT_LABEL.get(effort, effort)}</font>'

    def modules(self):
        from reportlab.platypus import KeepTogether, Spacer  # noqa: PLC0415

        mods = self.b.module_changes
        head, cmds = self._head(("MODULE", "TODAY", "TARGET", "CHANGE", "EFFORT"))
        data = [head]
        cmds.append(("LINEBELOW", (0, 1), (-1, -1), 0.5, self.C(RULE)))
        for r, m in enumerate(mods, start=1):
            first = [self.P(m.module, size=8.8, colour=INK, bold=True, leading=11)]
            if m.path and m.path != m.module:
                first.append(self.P(m.path, size=6.8, colour=MUTED))
            data.append([
                first, self._today_cell(m.current, m.current_status),
                self.P(m.target or "—", size=8.5, colour=INK, bold=True, leading=10.5),
                self.H(self.pill(CHANGE_LABEL[m.change_type], *CHANGE[m.change_type]), size=8) if m.change_type else "",
                self.H(self.effort(m.effort), size=8) if m.effort else "",
            ])
            cmds += self._today_cmds(1, r, m.current_status) + [("BACKGROUND", (2, r), (2, r), self.C(TARGET))]
        body: list = self.change_mix() + [self.table(data, [46, 36, 44, 24, 24], cmds, repeat=1)]
        detailed = [m for m in mods if _clean(m.changes)]
        if detailed:
            body += [Spacer(1, 8), self.P("What changes", size=9.5, colour=INK, bold=True, after=4)]
            for m in detailed:
                title = f'<font name="{self.bold}" color="#{INK}">{self.esc(m.module)}</font>&nbsp;&nbsp;'
                if m.change_type:
                    title += self.pill(CHANGE_LABEL[m.change_type], *CHANGE[m.change_type], size=7)
                if m.current or m.target:
                    title += (f'&nbsp;&nbsp;<font color="#{MUTED}" size="7.5">{self.esc(m.current or "—")} '
                              f'{self.g("→")} {self.esc(m.target or "—")}</font>')
                edge = CHANGE_BAR.get(m.change_type, "D4D4D4")
                card = self.table([[[self.H(title, size=9.5, after=3)] + self.bullets(m.changes, mark="•", mark_colour=ACCENT, size=8.6)]],
                                  [174], [("BACKGROUND", (0, 0), (-1, -1), self.C(PANEL)),
                                          ("LINEBEFORE", (0, 0), (0, -1), 2.5, self.C(edge)),
                                          ("LEFTPADDING", (0, 0), (-1, -1), 10), ("TOPPADDING", (0, 0), (-1, -1), 7),
                                          ("BOTTOMPADDING", (0, 0), (-1, -1), 7)])
                body += [KeepTogether([card]), Spacer(1, 5)]
        self.section("modules", body)

    def tradeoffs(self):
        items = self.b.trade_offs
        data = [[self.P("DECISION", size=6.8, colour=MUTED, bold=True),
                 self.P("WHAT WE GAIN", size=6.8, colour=GAIN[1], bold=True),
                 self.P("WHAT IT COSTS", size=6.8, colour=COST[1], bold=True)]]
        data += [[self.P(t.decision, size=8.8, colour=INK, bold=True, leading=11),
                  self.P(t.gain or "—", size=8.5, leading=11), self.P(t.cost or "—", size=8.5, leading=11)]
                 for t in items]
        self.section("tradeoffs", [self.table(data, [50, 62, 62], [
            ("BACKGROUND", (1, 1), (1, -1), self.C(GAIN[0])), ("BACKGROUND", (2, 1), (2, -1), self.C(COST[0])),
            ("LINEBELOW", (0, 1), (-1, -1), 0.5, self.C(WHITE)), ("VALIGN", (0, 1), (-1, -1), "TOP"),
        ], repeat=1)])

    def timeline(self):
        from reportlab.platypus import Image, Spacer  # noqa: PLC0415

        body: list = []
        chart = timeline_png(self.b)
        if chart:
            png, height_in = chart
            body += [Image(io.BytesIO(png), width=174 * self.mm, height=height_in * 72), Spacer(1, 4)]
        data = []
        for m in sorted_milestones(self.b):
            data.append([self.H(f'{self.glyph("●", _kind_dot(m.kind))}&nbsp;'
                                f'<font name="{self.bold}" color="#{INK}">{self.esc(pretty_date(m.date))}</font>', size=8.8),
                         self.P(m.label, size=8.8),
                         self.P(MILESTONE_LABEL.get(m.kind, "Milestone"), size=7.5, colour=_kind_ink(m.kind), bold=True)])
        body.append(self.table(data, [34, 110, 30], [("LINEBELOW", (0, 0), (-1, -1), 0.5, self.C(RULE))]))
        self.section("timeline", body)

    def constraints(self):
        self.section("constraints", self.bullets(self.b.constraints, mark="■", mark_colour=ACCENT))

    def success(self):
        from reportlab.platypus import Spacer  # noqa: PLC0415

        body: list = []
        if self.b.success_measures:
            data = [[self.P(h, size=6.8, colour=MUTED, bold=True) for h in ("MEASURE", "TODAY", "TARGET")]]
            data += [[self.P(m.metric, size=8.8, colour=INK, bold=True), self.P(_value(m.current), size=8.8, colour=MUTED),
                      self.P(m.target or "—", size=9.5, colour=ACCENT, bold=True)] for m in self.b.success_measures]
            body += [self.table(data, [80, 44, 50], [("BACKGROUND", (0, 0), (-1, 0), self.C(PANEL)),
                                                    ("BACKGROUND", (2, 1), (2, -1), self.C(TARGET)),
                                                    ("LINEBELOW", (0, 1), (-1, -1), 0.5, self.C(WHITE))], repeat=1),
                     Spacer(1, 6), self.P("Acceptance criteria", size=9.5, colour=INK, bold=True, after=3)]
        body += self.bullets(self.b.success_criteria, mark="○", mark_colour=ACCENT)
        self.section("success", body)

    def people(self):
        people = self.b.stakeholders
        rows = []
        for i in range(0, len(people), 2):
            pair = people[i:i + 2]
            cells = [[self.P(s.name, size=9.5, colour=INK, bold=True), self.P(s.role or "", size=8, colour=MUTED)]
                     for s in pair]
            rows.append([cells[0], "", cells[1] if len(cells) > 1 else ""])
        cmds = [("BACKGROUND", (0, r), (0, r), self.C(PANEL)) for r in range(len(rows))]
        cmds += [("BACKGROUND", (2, r), (2, r), self.C(PANEL)) for r, row in enumerate(rows) if row[2]]
        t = self.table(rows, [85, 4, 85], cmds + [("LINEBELOW", (0, 0), (-1, -2), 5, self.C(WHITE))])
        self.section("people", [t])

    def risks(self):
        cols = (("ASSUMPTIONS", self.b.assumptions, GREY_TINT, GREY_INK, MUTED, "None recorded."),
                ("RISKS", self.b.risks, RED_TINT, RED, RED, "None recorded."),
                ("OPEN QUESTIONS", self.b.open_questions, GREY_TINT, GREY_INK, MUTED, "None."))
        head, body, cmds = [], [], []
        for i, (label, items, tint, ink, mark, empty) in enumerate(cols):
            head += [self.P(label, size=6.8, colour=ink, bold=True)] + ([""] if i < 2 else [])
            body += [self.bullets(items, mark="•", mark_colour=mark, size=8.5, empty=empty)] + ([""] if i < 2 else [])
            cmds += [("BACKGROUND", (i * 2, 0), (i * 2, 0), self.C(tint)),
                     ("BACKGROUND", (i * 2, 1), (i * 2, 1), self.C(PANEL))]
        t = self.table([head, body], [56, 3, 56, 3, 56], cmds + [("VALIGN", (0, 1), (-1, 1), "TOP")])
        self.section("risks", [t])

    def repository(self):
        from reportlab.platypus import Spacer  # noqa: PLC0415

        text = repository_line(self.b)
        self.story += [Spacer(1, 12), self.H(f'<font name="{self.bold}" color="#{MUTED}" size="6.8">LEGACY REPOSITORY</font>'
                                            f"&nbsp;&nbsp;&nbsp;{self.esc(text)}", size=8)]

    def render(self, path: str) -> None:
        from reportlab.platypus import SimpleDocTemplate  # noqa: PLC0415

        mm = self.mm
        name = self.b.system_name or "Migration"

        def on_page(canvas, doc):
            canvas.saveState()
            # The page's top edge runs into PwC orange, like the Word document's strip.
            edge = canvas.beginPath()
            edge.rect(0, self.page[1] - 2.2 * mm, self.page[0], 2.2 * mm)
            canvas.clipPath(edge, stroke=0)
            canvas.linearGradient(0, 0, self.page[0], 0, [self.C(TINT_15), self.C(BRAND)], extend=False)
            canvas.restoreState()
            canvas.saveState()
            canvas.setFont(self.font, 7.5)
            canvas.setFillColor(self.C(MUTED))
            canvas.drawString(18 * mm, 9 * mm, f"{name} · Migration-intent brief")
            canvas.drawRightString(self.page[0] - 18 * mm, 9 * mm, f"Page {doc.page}")
            canvas.restoreState()
        doc = SimpleDocTemplate(path, pagesize=self.page, leftMargin=18 * mm, rightMargin=18 * mm,
                                topMargin=14 * mm, bottomMargin=16 * mm,
                                title=f"Migration-intent brief — {name}", author="SDLC Platform")
        self.title_band()
        self.facts()
        for key in _plan(self.b):
            getattr(self, key)()
        self.repository()
        doc.build(self.story, onFirstPage=on_page, onLaterPages=on_page)


# ── public API ───────────────────────────────────────────────────────────────


def render_brief(brief: MigrationIntentArtifact, path: str, meta: Optional[dict] = None) -> str:
    """Write the brief to `path` as .docx or .pdf (by extension). `meta` may carry the
    page version's `version` number and `status` for the title band."""
    ext = os.path.splitext(path)[1].lower()
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    if ext == ".docx":
        _Word(brief, meta or {}).render(path)
    elif ext == ".pdf":
        _Pdf(brief, meta or {}).render(path)
    else:
        raise ValueError(f"The brief renders as .docx or .pdf, not {ext or 'no extension'}")
    return path
