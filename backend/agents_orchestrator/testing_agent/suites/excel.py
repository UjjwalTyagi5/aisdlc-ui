"""The suite workbook — written for people, read back for runs.

A tester downloads the workbook, reviews it, may edit it, and the run executes what is in
it. So the format is both a document and an input:

- `About`      — what the suite is: test type, project, repository, branch, commit, what the
                 cases were derived from. The test type here is what `read_suite` trusts.
- `Test cases` — one row per case; columns are found BY HEADER, so reordering them in Excel
                 does not break a run.
- `Steps`      — FUNCTIONAL suites only: one row per step (Case ID, Step, Action, Target,
                 Value). The `Steps` column on `Test cases` is a readable summary; this
                 sheet is what runs.

Priority, scenario and step action carry drop-down validation, so an edit stays runnable.
"""
from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Any, Iterable, Optional, Union

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from agents_orchestrator.testing_agent.suites.models import (
    KIND_LABEL, STEP_ACTIONS, ApiCase, FunctionalCase, FunctionalStep, SuiteMeta, UnitCase, validate_cases,
)

HEAD_FILL = PatternFill("solid", fgColor="2D2D2D")
HEAD_FONT = Font(bold=True, color="FFFFFF")
ABOUT_KEY_FONT = Font(bold=True, color="5A5A5A")
TITLE_FONT = Font(bold=True, size=14, color="D04A02")
THIN = Side(style="thin", color="DDDDDD")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")

CASE_COLUMNS: dict[str, list[tuple[str, str, int]]] = {
    # (header, field, width)
    "unit": [("ID", "id", 10), ("Title", "title", 36), ("Module", "module", 30), ("Function", "function", 22),
             ("Scenario", "scenario", 12), ("Input", "input", 30), ("Expected result", "expected", 40),
             ("Priority", "priority", 10), ("Requirement", "requirement", 18)],
    "functional": [("ID", "id", 10), ("Title", "title", 36), ("Requirement", "requirement", 18),
                   ("Preconditions", "preconditions", 28), ("Steps", "steps", 60),
                   ("Expected result", "expected", 40), ("Priority", "priority", 10)],
    "api": [("ID", "id", 10), ("Title", "title", 34), ("Requirement", "requirement", 18), ("Method", "method", 9),
            ("Path", "path", 30), ("Headers", "headers", 22), ("Request body", "body", 34),
            ("Expected status", "expected_status", 10), ("Response must contain", "expected_body_contains", 30),
            ("Capture", "capture", 24), ("Priority", "priority", 10)],
}
STEP_COLUMNS = [("Case ID", 10), ("Step", 6), ("Action", 16), ("Target", 36), ("Value", 36)]
_ABOUT_ROWS = [("Test type", "kind"), ("Project", "project"), ("Source project", "source_project"),
               ("Repository", "repository"), ("Branch", "branch"),
               ("Commit", "commit"), ("Date", "generated_at"), ("Derived from", "sources"),
               ("Application notes", "app_notes")]


def _cell_value(field: str, case: Any) -> str:
    value = getattr(case, field, "")
    if field == "steps":
        return "\n".join(f"{i}. {s.describe()}" for i, s in enumerate(value, start=1))
    if field in ("headers", "capture"):
        return "; ".join(f"{k}={v}" for k, v in value.items())
    if field == "expected_body_contains":
        return "\n".join(value)
    return "" if value is None else str(value)


def _header_row(ws, headers: list[tuple[str, int]]) -> None:
    for col, (name, width) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, CELL_BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 22


def _validate_column(ws, header: str, headers: list[str], options: Iterable[str], rows: int) -> None:
    if header not in headers or rows < 1:
        return
    col = get_column_letter(headers.index(header) + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.error, dv.errorTitle = f"Choose one of: {', '.join(options)}", f"Invalid {header.lower()}"
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{max(rows + 1, 200)}")


def write_suite(meta: SuiteMeta, cases: list[Any], path: Optional[str] = None) -> bytes:
    """Write the workbook; returns its bytes (and writes `path` when given)."""
    wb = Workbook()
    about = wb.active
    about.title = "About"
    about["A1"] = f"{KIND_LABEL[meta.kind]} test cases"
    about["A1"].font = TITLE_FONT
    values = meta.model_dump()
    values["kind"] = KIND_LABEL[meta.kind]
    row = 3
    for label, key in _ABOUT_ROWS + [("Cases", "__count")]:
        about.cell(row=row, column=1, value=label).font = ABOUT_KEY_FONT
        about.cell(row=row, column=2, value=str(len(cases)) if key == "__count" else values.get(key, "")).alignment = WRAP_TOP
        row += 1
    about.cell(row=row + 1, column=1, value="Editing").font = ABOUT_KEY_FONT
    about.cell(row=row + 1, column=2, value=(
        "A run executes the cases in this workbook. Edit rows, add or remove cases, then upload it back to the "
        "project's Documents. Keep the column headers"
        + (" and the Steps sheet — it is what the browser follows." if meta.kind == "functional" else ".")
    )).alignment = WRAP_TOP
    about.column_dimensions["A"].width, about.column_dimensions["B"].width = 20, 90

    ws = wb.create_sheet("Test cases")
    columns = CASE_COLUMNS[meta.kind]
    _header_row(ws, [(h, w) for h, _f, w in columns])
    for r, case in enumerate(cases, start=2):
        for c, (_h, field, _w) in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=c, value=_cell_value(field, case))
            cell.alignment, cell.border = WRAP_TOP, CELL_BORDER
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(cases) + 1, 1)}"
    names = [h for h, _f, _w in columns]
    _validate_column(ws, "Priority", names, ("High", "Medium", "Low"), len(cases))
    _validate_column(ws, "Scenario", names, ("Happy path", "Error", "Edge"), len(cases))

    if meta.kind == "functional":
        steps = wb.create_sheet("Steps")
        _header_row(steps, STEP_COLUMNS)
        r = 2
        for case in cases:
            for n, step in enumerate(case.steps, start=1):
                for c, value in enumerate((case.id, n, step.action, step.target, step.value), start=1):
                    cell = steps.cell(row=r, column=c, value=value)
                    cell.alignment, cell.border = WRAP_TOP, CELL_BORDER
                r += 1
        _validate_column(steps, "Action", [h for h, _w in STEP_COLUMNS], STEP_ACTIONS, r - 2)

    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    if path:
        with open(path, "wb") as fh:
            fh.write(data)
    return data


def _rows(ws) -> list[dict[str, str]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [re.sub(r"\s+", " ", str(h or "")).strip().lower() for h in rows[0]]
    out = []
    for row in rows[1:]:
        if not any(v not in (None, "") for v in row):
            continue
        out.append({headers[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(row) if i < len(headers) and headers[i]})
    return out


class SuiteFormatError(ValueError):
    """The workbook is not a test case suite this platform can run — the message says why."""


def read_suite(source: Union[str, bytes]) -> tuple[SuiteMeta, list[Any], list[str]]:
    """(meta, cases, problems) from a suite workbook. Raises SuiteFormatError when it is not one."""
    try:
        wb = load_workbook(io.BytesIO(source) if isinstance(source, (bytes, bytearray)) else source, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise SuiteFormatError(f"the file is not a readable Excel workbook ({type(exc).__name__})") from exc
    if "About" not in wb.sheetnames or "Test cases" not in wb.sheetnames:
        raise SuiteFormatError("the workbook has no 'About' and 'Test cases' sheets — it is not a test case suite")

    about = {str(k or "").strip().lower(): ("" if v is None else str(v).strip())
             for k, v, *_ in wb["About"].iter_rows(min_row=3, max_col=2, values_only=True)}
    label = about.get("test type", "").lower()
    kind = next((k for k, v in KIND_LABEL.items() if v.lower() == label), None)
    if kind is None:
        raise SuiteFormatError(f"the About sheet names test type {label!r}; expected Unit, Functional or API")
    meta = SuiteMeta(kind=kind, **{key: about.get(lbl.lower(), "") for lbl, key in _ABOUT_ROWS if key != "kind"})

    by_header = {h.lower(): f for h, f, _w in CASE_COLUMNS[kind]}
    raw = []
    for row in _rows(wb["Test cases"]):
        item = {by_header[h]: v for h, v in row.items() if h in by_header}
        item.pop("steps", None)
        raw.append(item)

    if kind == "functional":
        if "Steps" not in wb.sheetnames:
            raise SuiteFormatError("a functional suite needs its 'Steps' sheet — it is what the browser follows")
        grouped: dict[str, list[tuple[float, dict]]] = {}
        for row in _rows(wb["Steps"]):
            cid = row.get("case id", "")
            try:
                order = float(row.get("step") or 0)
            except ValueError:
                order = 0
            grouped.setdefault(cid, []).append((order, {"action": row.get("action", ""), "target": row.get("target", ""), "value": row.get("value", "")}))
        for item in raw:
            item["steps"] = [s for _o, s in sorted(grouped.get(item.get("id", ""), []), key=lambda t: t[0])]

    cases, problems = validate_cases(kind, raw)
    return meta, cases, problems


def suite_filename(kind: str, repository: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9]+", "_", repository or "Project").strip("_") or "Project"
    return f"{stem}_{KIND_LABEL[kind]}_Test_Cases.xlsx"


def now_label() -> str:
    return datetime.now(timezone.utc).strftime("%d %b %Y %H:%M UTC")


__all__ = ["write_suite", "read_suite", "SuiteFormatError", "suite_filename", "now_label",
           "UnitCase", "FunctionalCase", "FunctionalStep", "ApiCase"]
