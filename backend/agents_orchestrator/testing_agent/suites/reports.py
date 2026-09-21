"""The run report — one workbook per run, every case accounted for.

A report that lists only what passed, or collapses "could not run" into "failed", is not
evidence anyone can sign off. Every case in the suite appears with exactly one status:

- Passed  — it ran and met its expected result
- Failed  — it ran and did not (the message says how)
- Error   — it could not be executed (a test that did not load, a page that did not open)
- Not run — nothing was executed for it (the message says why)

Sheets: `Summary` (what was run, against what, totals), `Results` (one row per case) and any
run-specific sheets (the generated test code, per-step outcomes, responses).
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from agents_orchestrator.testing_agent.suites.models import KIND_LABEL

STATUSES = ("Passed", "Failed", "Error", "Not run")
STATUS_FILL = {"Passed": "DCFCE7", "Failed": "FEE2E2", "Error": "FEF3C7", "Not run": "F1F5F9"}
HEAD_FILL = PatternFill("solid", fgColor="2D2D2D")
HEAD_FONT = Font(bold=True, color="FFFFFF")
KEY_FONT = Font(bold=True, color="5A5A5A")
TITLE_FONT = Font(bold=True, size=14, color="D04A02")
THIN = Side(style="thin", color="DDDDDD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
MAX_CELL = 32_000  # Excel's limit is 32,767 characters


@dataclass
class ResultRow:
    id: str
    title: str
    subject: str                     # what was exercised: module · function, a URL path, METHOD path
    status: str                      # one of STATUSES
    duration_ms: Optional[int] = None
    message: str = ""
    evidence: str = ""               # a screenshot file, a response excerpt, the test file


@dataclass
class ReportMeta:
    kind: str
    project: str = ""
    repository: str = ""
    branch: str = ""
    commit: str = ""
    suite_document: str = ""
    suite_status: str = ""           # the suite's approval state when it was run
    target_url: str = ""
    runner: str = ""
    started_at: str = ""
    duration_s: Optional[float] = None
    notes: list[str] = field(default_factory=list)


def totals(rows: list[ResultRow]) -> dict[str, int]:
    out = {s: 0 for s in STATUSES}
    for r in rows:
        out[r.status] = out.get(r.status, 0) + 1
    out["total"] = len(rows)
    return out


def verdict(rows: list[ResultRow]) -> str:
    t = totals(rows)
    if not rows:
        return "No cases"
    if t["Failed"] or t["Error"]:
        return "Failed"
    if t["Not run"]:
        return "Incomplete"
    return "Passed"


def _clip(text: Any) -> str:
    s = "" if text is None else str(text)
    return s if len(s) <= MAX_CELL else s[:MAX_CELL] + "\n… [truncated]"


def _sheet_table(ws, headers: list[tuple[str, int]], rows: list[list[Any]], status_col: Optional[int] = None) -> None:
    for c, (name, width) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        ws.column_dimensions[get_column_letter(c)].width = width
    for r, values in enumerate(rows, start=2):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=_clip(value))
            cell.alignment, cell.border = WRAP, BORDER
            if status_col is not None and c == status_col and value in STATUS_FILL:
                cell.fill = PatternFill("solid", fgColor=STATUS_FILL[value])
                cell.font = Font(bold=True)
    ws.freeze_panes = "B2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def write_report(meta: ReportMeta, rows: list[ResultRow], extra: Optional[dict[str, tuple[list[tuple[str, int]], list[list[Any]]]]] = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    t = totals(rows)
    ws["A1"] = f"{KIND_LABEL.get(meta.kind, meta.kind)} test report"
    ws["A1"].font = TITLE_FONT
    facts = [
        ("Result", verdict(rows)),
        ("Cases", str(t["total"])), ("Passed", str(t["Passed"])), ("Failed", str(t["Failed"])),
        ("Error", str(t["Error"])), ("Not run", str(t["Not run"])),
        ("Pass rate", f"{round(100 * t['Passed'] / t['total'])}%" if t["total"] else "—"),
        ("Project", meta.project), ("Repository", meta.repository), ("Branch", meta.branch),
        ("Commit tested", meta.commit), ("Test cases", f"{meta.suite_document} ({meta.suite_status})" if meta.suite_document else ""),
        ("Application URL", meta.target_url), ("Runner", meta.runner), ("Started", meta.started_at),
        ("Duration", f"{meta.duration_s:.0f} s" if meta.duration_s is not None else ""),
    ]
    row = 3
    for key, value in facts:
        if value in ("", None):
            continue
        ws.cell(row=row, column=1, value=key).font = KEY_FONT
        cell = ws.cell(row=row, column=2, value=value)
        cell.alignment = WRAP
        if key == "Result":
            cell.font = Font(bold=True, color="166534" if value == "Passed" else "991B1B")
        row += 1
    for note in meta.notes:
        ws.cell(row=row, column=1, value="Note").font = KEY_FONT
        ws.cell(row=row, column=2, value=note).alignment = WRAP
        row += 1
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 18, 100

    results = wb.create_sheet("Results")
    _sheet_table(results, [("ID", 10), ("Title", 36), ("Exercised", 34), ("Status", 10), ("Duration (ms)", 12),
                           ("Details", 60), ("Evidence", 30)],
                 [[r.id, r.title, r.subject, r.status, r.duration_ms, r.message, r.evidence] for r in rows], status_col=4)
    for name, (headers, values) in (extra or {}).items():
        _sheet_table(wb.create_sheet(name[:31]), headers, values)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v if v is not None else "").replace("|", "/")).strip() or "—"


def report_markdown(meta: ReportMeta, rows: list[ResultRow]) -> str:
    t = totals(rows)
    label = KIND_LABEL.get(meta.kind, meta.kind)
    lines = [f"# {meta.repository or meta.project} — {label} test report", "", "## Result", "",
             f"**{verdict(rows)}** — {t['Passed']} passed, {t['Failed']} failed, {t['Error']} could not run, "
             f"{t['Not run']} not run, of {t['total']} cases."]
    about = [("Branch", f"{meta.branch}" + (f" (commit {meta.commit[:7]})" if meta.commit else "")),
             ("Test cases", f"{meta.suite_document} ({meta.suite_status})" if meta.suite_document else ""),
             ("Application URL", meta.target_url), ("Runner", meta.runner), ("Started", meta.started_at),
             ("Duration", f"{meta.duration_s:.0f} s" if meta.duration_s is not None else "")]
    lines += ["", "## About this run", ""] + [f"- **{k}:** {v}" for k, v in about if v] + [f"- {n}" for n in meta.notes]
    lines += ["", "## Results", "", "| ID | Title | Exercised | Status | Details |", "|---|---|---|---|---|"]
    lines += [f"| {_cell(r.id)} | {_cell(r.title)} | {_cell(r.subject)} | {r.status} | {_cell(r.message)[:400]} |" for r in rows]
    return "\n".join(lines) + "\n"


def report_filename(kind: str, repository: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9]+", "_", repository or "Project").strip("_") or "Project"
    return f"{stem}_{KIND_LABEL[kind]}_Test_Report.xlsx"
