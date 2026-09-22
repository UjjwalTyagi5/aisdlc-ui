"""Read a stored `.docx` back as markdown, so a document without a page copy still opens.

A document an agent writes today keeps a page copy beside it (`artifact_page.py`) and the
app renders that. Documents written before page copies existed, documents from agents that
do not write the sibling markdown, and every hand-uploaded Word file have none — and the
page could then only offer a download, which is what "clicking it shows the file name
instead of the document" meant on the Design page.

This derives a readable view from the file itself. It is a PREVIEW, not the page copy: the
Word file keeps the title band, the palette and the fonts, and what comes back here is the
text, the headings, the lists and the tables in the order they appear.

TWO THINGS THE DESIGNED DOCUMENTS FORCED (architecture.docx, 15 tables, 51 paragraphs):

  * Every paragraph in them has style "Normal" — the design is direct run formatting, not
    named styles — so a heading is recognised by being bold and larger than the body text,
    with the level taken from how the sizes rank. Style names still win when there are any.
  * A merged cell is repeated by python-docx once per grid column the merge spans, so the
    title band's single cell came back 24 times in one row. Cells are de-duplicated by the
    element behind them, and a row left with one cell is a band, not a table.

Nothing here writes: no blob is added, no row is touched.
"""
from __future__ import annotations

import base64
import io
import logging
import re
from collections import Counter
from typing import Any, Optional

logger = logging.getLogger(__name__)

#: How much of a document's TEXT we are willing to render on a page. Figures do not count.
MAX_CHARS = 200_000
_STYLE_HEADING = re.compile(r"^Heading (\d)$", re.IGNORECASE)

#: A figure is carried into the page as a data URL — the formats a browser draws, up to a size
#: a page can hold. Anything else (an EMF, a very large scan) is named, and stays in the file.
MAX_FIGURE_BYTES = 4 * 1024 * 1024
MAX_FIGURES_BYTES = 16 * 1024 * 1024
_FIGURE_TYPES = {"image/png", "image/jpeg", "image/gif", "image/webp"}
_FIGURE_NOT_SHOWN = "_A figure here is only in the Word file._"


def _figures(element: Any, document: Any, budget: list[int]) -> list[str]:
    """The pictures inside a paragraph or table, in order, as markdown images.

    DIAGRAMS WERE DROPPED. The Design agent's documents draw their C4, sequence and ER
    diagrams as pictures (`WordCanvas.figure`), and a document read back without a page copy
    lost every one of them — the reader kept text only, so the page showed the captions
    ("Figure 3 · …") with nothing above them. `budget` is the bytes left for the document.
    """
    out = []
    for rid in element.xpath(".//a:blip/@r:embed"):
        part = document.part.related_parts.get(rid)
        blob = getattr(part, "blob", None)
        content_type = getattr(part, "content_type", "")
        if not blob or content_type not in _FIGURE_TYPES or len(blob) > MAX_FIGURE_BYTES or len(blob) > budget[0]:
            out.append(_FIGURE_NOT_SHOWN)
            continue
        budget[0] -= len(blob)
        out.append(f"![Figure](data:{content_type};base64,{base64.b64encode(blob).decode('ascii')})")
    return out


def can_preview(filename: str) -> bool:
    """Only Word documents. A spreadsheet or a deck is not text, and guessing is worse
    than the download the caller already offers."""
    return filename.lower().endswith(".docx")


def _cells(row: Any) -> list[str]:
    """A row's cells, each once — a merge repeats the same cell across the columns it spans."""
    out, seen = [], set()
    for cell in row.cells:
        key = id(cell._tc)
        if key in seen:
            continue
        seen.add(key)
        text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
        out.append(text.replace("|", "\\|").replace("\n", " "))
    return out


def _body_size(document: Any) -> float:
    """The size most of the text is set in — the baseline a heading stands above."""
    sizes: Counter[float] = Counter()
    for para in document.paragraphs:
        if not para.text.strip():
            continue
        for run in para.runs:
            if run.font.size is not None and run.text.strip():
                sizes[round(run.font.size.pt, 1)] += len(run.text)
    return sizes.most_common(1)[0][0] if sizes else 11.0


def _heading_sizes(document: Any, body: float) -> list[float]:
    """Distinct sizes used for bold text bigger than the body, largest first."""
    sizes = set()
    for para in document.paragraphs:
        if not para.text.strip():
            continue
        runs = [r for r in para.runs if r.text.strip()]
        if runs and all(r.bold for r in runs):
            pts = [round(r.font.size.pt, 1) for r in runs if r.font.size is not None]
            if pts and min(pts) > body:
                sizes.add(max(pts))
    return sorted(sizes, reverse=True)[:3]


def _paragraph(para: Any, body: float, heading_sizes: list[float]) -> Optional[str]:
    text = para.text.strip()
    if not text:
        return None
    style = (para.style.name if para.style is not None else "") or ""
    match = _STYLE_HEADING.match(style)
    if match:
        return f"{'#' * min(int(match.group(1)), 6)} {text}"
    if style.lower() == "title":
        return f"# {text}"
    if style.startswith("List Bullet"):
        return f"- {text}"
    if style.startswith("List Number"):
        return f"1. {text}"
    if style.lower() in ("quote", "intense quote"):
        return f"> {text}"

    runs = [r for r in para.runs if r.text.strip()]
    if runs and all(r.bold for r in runs):
        pts = [round(r.font.size.pt, 1) for r in runs if r.font.size is not None]
        size = max(pts) if pts else body
        if size in heading_sizes:
            return f"{'#' * (heading_sizes.index(size) + 2)} {text}"
        if size > body:
            return f"#### {text}"
        return f"**{text}**"
    return text


def _table(table: Any) -> list[str]:
    rows = [c for c in (_cells(r) for r in table.rows) if any(c)]
    if not rows:
        return []
    # A one-cell row is a band or a callout, not tabular data.
    if max(len(r) for r in rows) == 1:
        return [r[0] for r in rows]
    width = max(len(r) for r in rows)
    pad = lambda r: r + [""] * (width - len(r))  # noqa: E731
    lines = ["| " + " | ".join(pad(rows[0])) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines += ["| " + " | ".join(pad(r)) + " |" for r in rows[1:]]
    return ["\n".join(lines)]


def docx_markdown(data: bytes) -> Optional[str]:
    """The document's text as markdown, or None when it cannot be read as one."""
    try:
        from docx import Document  # noqa: PLC0415
        from docx.table import Table  # noqa: PLC0415
        from docx.text.paragraph import Paragraph  # noqa: PLC0415
    except ImportError:  # pragma: no cover — python-docx ships with the agents
        logger.warning("python-docx is not installed; no document preview")
        return None

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 — a corrupt or non-Word file is not an error here
        logger.info("document preview: not readable as .docx (%s)", type(exc).__name__)
        return None

    body = _body_size(document)
    headings = _heading_sizes(document, body)
    budget = [MAX_FIGURES_BYTES]
    blocks: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            blocks.extend(_figures(child, document, budget))
            block = _paragraph(Paragraph(child, document), body, headings)
            if block and block != (blocks[-1] if blocks else None):
                blocks.append(block)
        elif tag == "tbl":
            blocks.extend(_table(Table(child, document)))
            blocks.extend(_figures(child, document, budget))

    # The text limit counts text: a figure is one line of markdown and many bytes.
    kept: list[str] = []
    room = MAX_CHARS
    for block in blocks:
        if block.startswith("![Figure](data:"):
            kept.append(block)
            continue
        if len(block) > room:
            head = block[:room].rsplit("\n", 1)[0] if room > 0 else ""
            if head.strip():
                kept.append(head)
            kept.append("_The rest of this document is in the Word file._")
            break
        kept.append(block)
        room -= len(block) + 2
    markdown = "\n\n".join(kept).strip()
    return markdown or None
