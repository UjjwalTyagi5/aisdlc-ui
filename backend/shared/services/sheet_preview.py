"""Read a stored Excel workbook back as sheets of rows, so it opens on the page like a document.

Agents file spreadsheets as well as Word documents — the Project Manager's estimates and plans,
the Testing agent's case suites and run reports, anything `export_document` writes as .xlsx.
They went through approval like any document, but opening one offered only a download: the
page's viewer rendered markdown, and a workbook has none. A CSV file is read the same way, as a
workbook of one sheet.

This returns the workbook as data — every sheet, in order, as rows of display strings — for the
page to lay out as a grid. It is a VIEW of the file: formulas show their last computed value,
number formats are applied loosely, and a very large sheet is cut and says so. The Excel file
is still the document; the page offers it for download beside the view.

Nothing here writes.
"""
from __future__ import annotations

import csv
import io
import logging
from datetime import date, datetime, time
from typing import Any, Optional

logger = logging.getLogger(__name__)

#: What one sheet may show on a page. Past these the view says it was cut.
MAX_ROWS = 1000
MAX_COLS = 60
MAX_SHEETS = 20
MAX_CELL = 2000

_SHEET_EXTS = (".xlsx", ".xlsm")
_CSV_EXTS = (".csv", ".tsv")


def can_preview_sheet(filename: str) -> bool:
    return filename.lower().endswith(_SHEET_EXTS)


def can_preview_csv(filename: str) -> bool:
    return filename.lower().endswith(_CSV_EXTS)


def _display(value: Any) -> str:
    """A cell as a person reads it. Dates without a time lose the midnight; whole floats
    lose the `.0` Excel stores every number with."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0, 0) else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.6g}"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    text = str(value)
    return text if len(text) <= MAX_CELL else text[:MAX_CELL] + "…"


def xlsx_sheets(data: bytes) -> Optional[list[dict]]:
    """[{name, rows: [[str]], rows_total, cols_total, truncated}], or None when the bytes are
    not a readable workbook. Trailing empty rows and columns are dropped."""
    try:
        from openpyxl import load_workbook  # noqa: PLC0415
    except ImportError:  # pragma: no cover — openpyxl ships with the agents
        logger.warning("openpyxl is not installed; no spreadsheet preview")
        return None
    try:
        # read_only streams rows; data_only shows each formula's cached value, which is what
        # Excel itself displays — the formula text would read as noise on a page.
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — a corrupt or non-Excel file is not an error here
        logger.info("sheet preview: not readable as a workbook (%s)", type(exc).__name__)
        return None

    try:
        return [_sheet(ws.title, ws.iter_rows(values_only=True), getattr(ws, "max_column", 0) or 0)
                for ws in workbook.worksheets[:MAX_SHEETS]]
    finally:
        workbook.close()


def csv_sheets(data: bytes, filename: str) -> Optional[list[dict]]:
    """A CSV (or TSV) file as a workbook of one sheet named after the file, or None when the
    bytes are not text. The delimiter is sniffed, so a semicolon-separated export reads too."""
    text = decode_text(data)
    if text is None:
        return None
    if filename.lower().endswith(".tsv"):
        dialect: Any = csv.excel_tab
    else:
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
    name = filename.rsplit("/", 1)[-1].rsplit(".", 1)[0] or "Sheet1"
    return [_sheet(name, csv.reader(io.StringIO(text), dialect))]


def decode_text(data: bytes) -> Optional[str]:
    """UTF-8 (with or without Excel's byte-order mark), else Windows-1252 — the two a CSV export
    is written in. NUL bytes mean it is not text at all."""
    if b"\x00" in data[:4096]:
        return None
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _sheet(name: str, raw_rows: Any, width_hint: int = 0) -> dict:
    """One sheet: rows of display strings, trailing empty rows and columns dropped, every row
    as wide as the widest so the grid has straight columns, cut past MAX_ROWS/MAX_COLS."""
    rows: list[list[str]] = []
    rows_total = 0
    widest = 0
    for raw in raw_rows:
        rows_total += 1
        if len(rows) >= MAX_ROWS:
            continue            # keep counting, stop keeping
        full = [_display(v) for v in (raw or ())]
        while full and full[-1] == "":
            full.pop()
        width_hint = max(width_hint, len(full))
        cells = full[:MAX_COLS]
        rows.append(cells)
        widest = max(widest, len(cells))
    while rows and not any(rows[-1]):
        rows.pop()
    rows = [r + [""] * (widest - len(r)) for r in rows]
    cols_total = max(widest, width_hint)
    return {
        "name": name,
        "rows": rows,
        "rows_total": rows_total,
        "cols_total": cols_total,
        "truncated": rows_total > MAX_ROWS or cols_total > MAX_COLS,
    }
