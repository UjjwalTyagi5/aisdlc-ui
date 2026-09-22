"""Read a stored PowerPoint deck back as markdown, slide by slide, so a deck opens on the page.

The Design agent files decks (the architecture presentation) beside its Word documents. They
went through approval like any document, but opening one offered only a download: a browser
cannot draw a .pptx, and the page's viewer rendered markdown.

This derives a readable outline from the file: each slide's title as a heading, its text in
reading order (top to bottom, then left to right), its tables, and the speaker notes. It is a
PREVIEW, not the deck — layout, pictures and colours stay in the file, which the page offers
beside the view, and the view says it was read from the slides.

Nothing here writes.
"""
from __future__ import annotations

import io
import logging
from typing import Any, Iterator, Optional

logger = logging.getLogger(__name__)

#: How much of a deck we are willing to render on a page.
MAX_CHARS = 200_000


def can_preview_slides(filename: str) -> bool:
    return filename.lower().endswith(".pptx")


def _shapes(shapes: Any) -> Iterator[Any]:
    """Every shape, groups opened, in the order a person reads a slide."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: PLC0415

    ordered = sorted(shapes, key=lambda s: ((s.top or 0), (s.left or 0)))
    for shape in ordered:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes)
        else:
            yield shape


def _text(shape: Any) -> Optional[str]:
    """A text box as a paragraph, or as a list when it holds several — which on a slide is
    what a text box almost always is."""
    paragraphs = [(p.level, "".join(r.text for r in p.runs).strip()) for p in shape.text_frame.paragraphs]
    paragraphs = [(level, text) for level, text in paragraphs if text]
    if not paragraphs:
        return None
    if len(paragraphs) == 1:
        return paragraphs[0][1]
    return "\n".join(f"{'  ' * level}- {text}" for level, text in paragraphs)


def _table(shape: Any) -> Optional[str]:
    rows = []
    for row in shape.table.rows:
        cells = ["" if cell.is_spanned else cell.text.strip().replace("|", "\\|").replace("\n", " ")
                 for cell in row.cells]
        if any(cells):
            rows.append(cells)
    if not rows:
        return None
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def pptx_markdown(data: bytes) -> Optional[str]:
    """The deck's slides as markdown, or None when the bytes are not a readable deck."""
    try:
        from pptx import Presentation  # noqa: PLC0415
    except ImportError:  # pragma: no cover — python-pptx ships with the agents
        logger.warning("python-pptx is not installed; no deck preview")
        return None
    try:
        deck = Presentation(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 — a corrupt or non-deck file is not an error here
        logger.info("deck preview: not readable as .pptx (%s)", type(exc).__name__)
        return None

    blocks: list[str] = []
    for number, slide in enumerate(deck.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        blocks.append(f"## Slide {number}" + (f" · {title}" if title else ""))
        for shape in _shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            block = None
            if getattr(shape, "has_table", False) and shape.has_table:
                block = _table(shape)
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                block = _text(shape)
            if block:
                blocks.append(block)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                blocks.append("> **Speaker notes:** " + " ".join(notes.split()))

    markdown = "\n\n".join(blocks).strip()
    if not markdown:
        return None
    if len(markdown) > MAX_CHARS:
        markdown = (markdown[:MAX_CHARS].rsplit("\n", 1)[0]
                    + "\n\n_The rest of this deck is in the PowerPoint file._")
    return markdown
